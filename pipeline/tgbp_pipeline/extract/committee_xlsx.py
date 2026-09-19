"""T-108: เอกสาร กมธ.ติดตามงบ — Excel/XLS (A5, 03-DATA-PIPELINE.md §4.4)

Generic table detector ต่อ sheet:
1. หาแถว "header" ที่มี >= `MIN_HEADER_TEXT_CELLS` เซลล์ข้อความ (ไม่ใช่ตัวเลขล้วน) แล้วภายใน
   `HEADER_LOOKAHEAD` แถวถัดไปมีแถวที่มีตัวเลขอย่างน้อย 1 เซลล์ (แถวระหว่างกลางถือเป็น sub-header/
   unit-hint เช่น `(บาท)`) → ได้ "ตาราง" หนึ่งก้อน (header, data_start, data_end)
2. ถ้าหลาย candidate ใน sheet เดียว เลือกก้อนที่ใหญ่ที่สุด (จำนวนแถวข้อมูลมากสุด) — ไม่พยายามหา
   หลายตารางต่อ sheet (03 §4.4: "ไม่จำเป็น — เอาตารางแรกที่ใหญ่สุด")
3. map header → canonical field ด้วย synonym dict + rapidfuzz; ต้องได้อย่างน้อย `item_name` +
   `amount_thb` และ**ต้องรู้หน่วยเงิน**จาก header/แถว title/sub-header (บาท/พันบาท/ล้านบาท) —
   ไม่รู้หน่วย = **ห้ามเดา** (CLAUDE.md N3) → ไม่ใส่ budget_lines
4. map ได้ → เขียน parquet `.cache/committee/{doc_id}.parquet` (dataset `committee_table`)
   map ไม่ได้ (ไม่พบตาราง, ขาด field หลัก, หรือไม่รู้หน่วย) → เก็บเป็น `DocChunk` แบบตาราง
   (`extract.office_text.DocChunk`) ไปที่ `.cache/docs/{doc_id}.json.gz` แทน — ไม่ใส่ budget_lines

`.XLS` ที่หัวไฟล์เป็น zip (`PK`) จริง ๆ คือ `.xlsx` ที่ตั้งนามสกุลผิด (พบจริงใน
`องค์การโคนม/BI*.XLS` ทั้ง 13 ไฟล์ 19 ก.ย. 2569) → ตรวจ magic bytes ก่อนเลือก reader เสมอ
ไม่ใช่เชื่อนามสกุลไฟล์ตรง ๆ; ถ้า xlrd เปิดไม่ได้จริง (`.xls` แท้ที่เสีย) → skip + เหตุผล
(เครื่องนี้ไม่มี LibreOffice — ใช้ `soffice` เฉพาะถ้าเจอใน PATH)

**แก้ 19 ก.ย. 2569 (รอบตรวจซ้ำ):** แค่ route ไปหา `_iter_openpyxl_sheets` ตาม magic bytes
ไม่พอ — `openpyxl.load_workbook(str(path), ...)` **เอง** ก็เช็คนามสกุลไฟล์จาก path string
ที่ส่งเข้าไปด้วย แล้ว raise `"does not support the old .xls file format"` ทั้งที่เนื้อไฟล์เป็น
zip/xlsx จริง (ยืนยันจริงกับทั้ง 13 ไฟล์ `BI*.XLS`) → ต้องเปิดเป็น **binary file object**
(`open(path, "rb")`) แล้วส่ง file object นั้นให้ `load_workbook` แทนเสมอ ไม่ใช่ path string
"""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import xlrd
from rapidfuzz import fuzz

from tgbp_pipeline.config import PipelineConfig
from tgbp_pipeline.extract.office_text import Atom, DocChunk, chunk_atoms, write_doc_chunks_gz
from tgbp_pipeline.normalize import money
from tgbp_pipeline.normalize.thai_text import clean, nfc_only
from tgbp_pipeline.util.hash import doc_id_for_path
from tgbp_pipeline.util.hash import source_id as compute_source_id

DATASET = "committee_table"
COMMITTEE_ROOT = "กมธ.ติดตามงบ"
COMMITTEE_CACHE_DIRNAME = "committee"

MIN_HEADER_TEXT_CELLS = 4
HEADER_LOOKAHEAD = 3
TITLE_LOOKBACK = 6

# canonical field -> คำพ้อง (ผ่าน thai_text.clean แล้ว) ที่ใช้จับคู่กับ header cell — 03 §4.4
CANONICAL_SYNONYMS: dict[str, tuple[str, ...]] = {
    "item_name": (
        "ชื่อโครงการ",
        "รายการ",
        "ชื่อรหัสงบประมาณ",
        "แผนงาน/ชื่อโครงการ",
        "โครงการ",
        "งาน/โครงการ",
    ),
    "amount_thb": (
        "งบประมาณ",
        "งบประมาณที่อนุมัติ",
        "วงเงินที่ได้รับอนุมัติ",
        "เงินงบประมาณ",
        "วงเงินงบประมาณที่ขอรับจัดสรร",
        "จำนวนเงิน",
        "วงเงิน",
        "เงินที่ได้รับอนุมัติ",
    ),
    "agency": (
        "หน่วยงาน",
        "ผู้ได้รับการสนับสนุน",
        "หน่วยงานที่ได้รับการสนับสนุน",
        "กรม",
        "ส่วนราชการ",
    ),
    "plan": (
        "แผนงาน",
        "ยุทธศาสตร์",
        "ยุทธศาสตร์/แผนงาน",
    ),
    "province": ("จังหวัด",),
    "item_qty": ("จำนวน", "ปริมาณ"),
}
# ลำดับความสำคัญตอน greedy-assign คอลัมน์ (item/amount ก่อน กัน field รองแย่งคอลัมน์ที่ควรเป็นของหลัก)
_FIELD_PRIORITY: tuple[str, ...] = (
    "item_name",
    "amount_thb",
    "agency",
    "plan",
    "province",
    "item_qty",
)
FUZZY_COLUMN_THRESHOLD = 85.0
REQUIRED_FIELDS: tuple[str, ...] = ("item_name", "amount_thb")

# ลำดับสำคัญ: ตรวจ "ล้านบาท"/"พันบาท" ก่อน "บาท" เพราะ "บาท" เป็น substring ของทั้งคู่
UNIT_MULTIPLIERS: tuple[tuple[str, int], ...] = (
    ("ล้านบาท", 1_000_000),
    ("พันบาท", 1_000),
    ("บาท", 1),
)

# หน่วยงานที่เป็นรัฐวิสาหกิจจริง (พบในโฟลเดอร์ กมธ.ติดตามงบ — 02 §A5); อื่น ๆ ถือเป็น central
# (กรม/กองทุน/เงินทุนหมุนเวียนของราชการ ไม่ใช่ อปท. จึงไม่ใช่ "local")
_STATE_ENTERPRISE_AGENCY_KEYWORDS: tuple[str, ...] = ("องค์การคลังสินค้า", "องค์การโคนม")

# อปท. (03 §5.3 ใช้เกณฑ์เดียวกันสำหรับ org_master): เทศบาล/องค์การบริหารส่วนจังหวัด-ตำบล/
# กรุงเทพมหานคร/เมืองพัทยา — ชื่อหน่วยงานที่ขึ้นต้นด้วยคำเหล่านี้ต้องเป็น gov_level="local"
# (บั๊กจริงที่แก้ 19 ก.ย. 2569: `กองทุนอนุรักษ์พลังงาน` 2,549/3,328 แถวเป็น อปท. แต่เดิมได้
# gov_level="central" หมดทุกแถว เพราะคำนวณจาก fallback_agency ของ sheet ครั้งเดียว ไม่ใช่ต่อแถว)
_LOCAL_GOV_AGENCY_PREFIXES: tuple[str, ...] = (
    "เทศบาล",
    "องค์การบริหารส่วน",
    "กรุงเทพมหานคร",
    "เมืองพัทยา",
)
_MINISTRY_LABEL_RE = re.compile(r"กระทรวง\s*[:：]\s*(\S.*)")
_FISCAL_YEAR_4_DIGIT_RE = re.compile(r"25\d{2}")
_FISCAL_YEAR_2_DIGIT_RE = re.compile(
    r"(?:ปีงบประมาณ|ปีงบ|ของบประมาณ|ของบฯ|ของบ|งบประมาณ|งบฯ|งบ|พ\.ศ\.?|ปี)[\s.:]{0,4}(\d{2})(?!\d)"
)
_FISCAL_YEAR_MIN_BE = 2540
_FISCAL_YEAR_MAX_BE = 2600


class UnsupportedWorkbookError(ValueError):
    """เปิดไฟล์ xlsx/xls ไม่ได้จริง (เสีย/รูปแบบไม่รองรับ) — ต้อง skip + เหตุผล"""


# -- table region detection -------------------------------------------------


def _is_blank_cell(value: object) -> bool:
    return value is None or (isinstance(value, str) and clean(value) == "")


def _is_number_cell(value: object) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, int | float):
        return True
    if isinstance(value, str):
        text = clean(value)
        if text in ("", "-"):
            return False
        return money.parse_baht(text) is not None
    return False


def _is_header_text_cell(value: object) -> bool:
    """เซลล์ที่นับเป็น "ข้อความ header" — string ที่ไม่ใช่ตัวเลขล้วนและไม่ว่าง"""
    if not isinstance(value, str):
        return False
    text = clean(value)
    if text == "":
        return False
    return money.parse_baht(text) is None


@dataclass(frozen=True)
class TableRegion:
    header_row_idx: int
    header: tuple
    data_start_idx: int
    data_end_idx: int  # exclusive


def _find_data_start(rows: list[tuple], header_idx: int) -> int | None:
    limit = min(header_idx + 1 + HEADER_LOOKAHEAD, len(rows))
    for j in range(header_idx + 1, limit):
        if any(_is_number_cell(v) for v in rows[j]):
            return j
    return None


def _find_table_end(rows: list[tuple], data_start: int) -> int:
    """แถวสุดท้ายของตาราง (exclusive) — หยุดเมื่อเจอแถวว่างล้วน 2 แถวติดกัน หรือจบ sheet"""
    end = data_start
    blank_streak = 0
    for i in range(data_start, len(rows)):
        if all(_is_blank_cell(v) for v in rows[i]):
            blank_streak += 1
            if blank_streak >= 2:
                break
        else:
            blank_streak = 0
            end = i + 1
    return end


def detect_table_region(rows: list[tuple]) -> TableRegion | None:
    """หาก้อนตารางที่ "ใหญ่สุด" ใน sheet เดียว (03 §4.4) — คืน `None` ถ้าไม่พบเลย"""
    candidates: list[tuple[int, int, int]] = []
    for i, row in enumerate(rows):
        text_count = sum(1 for v in row if _is_header_text_cell(v))
        if text_count < MIN_HEADER_TEXT_CELLS:
            continue
        data_start = _find_data_start(rows, i)
        if data_start is None:
            continue
        data_end = _find_table_end(rows, data_start)
        if data_end <= data_start:
            continue
        candidates.append((i, data_start, data_end))

    if not candidates:
        return None

    header_idx, data_start, data_end = max(candidates, key=lambda c: (c[2] - c[1], -c[0]))
    return TableRegion(
        header_row_idx=header_idx,
        header=rows[header_idx],
        data_start_idx=data_start,
        data_end_idx=data_end,
    )


def detect_unit(rows: list[tuple], region: TableRegion) -> tuple[str, int] | None:
    """หาหน่วยเงินจากแถว title เหนือ header + header เอง + sub-header/unit-hint ระหว่าง header กับข้อมูล

    ตาม CLAUDE.md N3: ไม่พบคำว่า บาท/พันบาท/ล้านบาท เลย → คืน `None` (ห้ามเดาหน่วย)
    """
    start = max(0, region.header_row_idx - TITLE_LOOKBACK)
    end = region.data_start_idx  # รวมแถว title, header, sub-header (ไม่รวมแถวข้อมูลตัวเลขจริง)
    texts: list[str] = []
    for idx in range(start, end):
        for value in rows[idx]:
            if isinstance(value, str):
                texts.append(clean(value))
    haystack = " | ".join(texts)
    for keyword, multiplier in UNIT_MULTIPLIERS:
        if keyword in haystack:
            return keyword, multiplier
    return None


def map_columns(header: tuple) -> dict[str, int]:
    """คืน `{canonical_field: column_index}` — greedy assign ตามลำดับ `_FIELD_PRIORITY`"""
    header_texts = [clean(str(v)) if v is not None else "" for v in header]
    assigned: dict[str, int] = {}
    used_cols: set[int] = set()

    for field_name in _FIELD_PRIORITY:
        synonyms = CANONICAL_SYNONYMS[field_name]
        best_col: int | None = None
        best_score = 0.0
        for col_idx, text in enumerate(header_texts):
            if col_idx in used_cols or not text:
                continue
            score = max((100.0 if text == syn else fuzz.WRatio(text, syn)) for syn in synonyms)
            if score > best_score:
                best_score = score
                best_col = col_idx
        if best_col is not None and best_score >= FUZZY_COLUMN_THRESHOLD:
            assigned[field_name] = best_col
            used_cols.add(best_col)

    return assigned


# -- gov_level / ministry / fiscal year (จากเนื้อไฟล์ ไม่ใช่ item_parser/org_master ของ T-102..104) --


_NUMBERED_SUBFOLDER_RE = re.compile(r"^\d+[.)]\s")


def _agency_from_rel_path(rel_path: str) -> str | None:
    """`กมธ.ติดตามงบ/ครั้งที่ N (...)/<เรื่อง>/[<หน่วยงาน>/]ไฟล์` (02 §C, 02 §A5)

    ข้าม subfolder ที่เป็นป้ายจัดหมวดข้อมูลทั่วไป (ขึ้นต้นด้วยเลข+จุด/วงเล็บ+เว้นวรรค เช่น
    `1. ข้อมูลโครงการ`) เพราะไม่ใช่ชื่อหน่วยงาน — พบจริงใน
    `กองทุนอนุรักษ์พลังงาน/1. ข้อมูลโครงการ/รวมข้อมูลโครงการ...xlsx` (มี subfolder เกินมา
    1 ชั้นเทียบกับ 02 §C ปกติ) ถ้าตัดออกแล้วไม่เหลือเลย ใช้ list เดิมทั้งหมดแทน (safety net)
    """
    parts = rel_path.split("/")
    middle = [clean(p) for p in parts[1:-1]]  # ตัด root folder และชื่อไฟล์
    named = [p for p in middle if not _NUMBERED_SUBFOLDER_RE.match(p)]
    candidates = named or middle
    if not candidates:
        return None
    return candidates[-1]


def _gov_level_for_agency(agency: str | None, extra_text: str) -> str:
    """ต้องเรียกต่อแถวด้วย `agency` ที่ resolve แล้วจริงของแถวนั้น (ห้ามคำนวณครั้งเดียวต่อ sheet
    ด้วย fallback agency — ดูคอมเมนต์ `_LOCAL_GOV_AGENCY_PREFIXES`)
    """
    if "รัฐวิสาหกิจ" in extra_text:
        return "state_enterprise"
    if agency:
        if any(agency.startswith(prefix) for prefix in _LOCAL_GOV_AGENCY_PREFIXES):
            return "local"
        if any(kw in agency for kw in _STATE_ENTERPRISE_AGENCY_KEYWORDS):
            return "state_enterprise"
    return "central"


def _detect_ministry_label(rows: list[tuple], region: TableRegion) -> str | None:
    start = max(0, region.header_row_idx - TITLE_LOOKBACK)
    for idx in range(start, region.data_start_idx):
        for value in rows[idx]:
            if not isinstance(value, str):
                continue
            match = _MINISTRY_LABEL_RE.search(clean(value))
            if match:
                label = match.group(1).strip()
                # ตัด placeholder ของ template ที่ยังไม่ถูกแทนค่าจริง (เช่น `chk_min_name`)
                if label and not label.startswith("chk_"):
                    return label
    return None


_AS_OF_DATE_RE = re.compile(r"ณ\s*วันที่")
_SHEET_NAME_LEADING_YEAR_RE = re.compile(r"^(\d{2})(?!\d)")


def _year_candidates_in_text(text: str) -> set[int]:
    found: set[int] = set()
    for m in _FISCAL_YEAR_4_DIGIT_RE.finditer(text):
        year = int(m.group(0))
        if _FISCAL_YEAR_MIN_BE <= year <= _FISCAL_YEAR_MAX_BE:
            found.add(year)
    if not found:
        for m in _FISCAL_YEAR_2_DIGIT_RE.finditer(text):
            year = 2500 + int(m.group(1))
            if _FISCAL_YEAR_MIN_BE <= year <= _FISCAL_YEAR_MAX_BE:
                found.add(year)
    return found


def _fiscal_year_from_sheet_name(sheet_name: str) -> int | None:
    """ปีงบจากชื่อ sheet เอง — แหล่งที่เชื่อถือได้สุด (สั้น ตรงประเด็น ไม่ปนวันที่ปรับปรุงข้อมูล)

    รองรับทั้ง `"ปี 63"`/`"ปี 64 กลุ่ม 7"` (มีคำ "ปี" นำ ใช้ regex ปกติ) และ `"61เพิ่มเติม"`
    (ไม่มีคำนำ เลข 2 หลักขึ้นต้นชื่อ sheet ตรง ๆ — พบจริงใน กองทุนอนุรักษ์พลังงาน)
    """
    text = clean(sheet_name)
    candidates = _year_candidates_in_text(text)
    if not candidates:
        m = _SHEET_NAME_LEADING_YEAR_RE.match(text)
        if m:
            year = 2500 + int(m.group(1))
            if _FISCAL_YEAR_MIN_BE <= year <= _FISCAL_YEAR_MAX_BE:
                candidates.add(year)
    if len(candidates) == 1:
        return next(iter(candidates))
    return None


def _extract_fiscal_year_be(sheet_name: str, texts: list[str]) -> int | None:
    """ปีงบของทั้ง sheet — ลองจากชื่อ sheet ก่อนเสมอ (แม่นสุด) แล้วค่อย fallback ไปข้อความ
    title/header รอบตาราง โดยตัดข้อความ `"ณ วันที่ ..."` (วันที่ปรับปรุงข้อมูล ไม่ใช่ปีงบ) ทิ้ง
    ก่อนเสมอ

    **บั๊กจริงที่แก้ 19 ก.ย. 2569:** เดิมค้นหาปีจากทุก text ปนกันโดยไม่แยกลำดับความสำคัญ —
    sheet "ปี 63" ของ กองทุนอนุรักษ์พลังงาน มีทั้ง `"ปีงบประมาณ พ.ศ. 2563"` (ปีงบจริง) และ
    `"ข้อมูล ณ วันที่ 30 มิถุนายน 2567"` (วันที่ปรับปรุงข้อมูล) อยู่ในไตเติลเดียวกัน → เจอ 2 ปี
    ต่างกัน → กำกวม → คืน `None` ทุกแถว (`fiscal_year_be` เป็น null 100% ในทุก sheet ของไฟล์นี้)
    """
    from_sheet_name = _fiscal_year_from_sheet_name(sheet_name)
    if from_sheet_name is not None:
        return from_sheet_name

    filtered_texts = [t for t in texts if not _AS_OF_DATE_RE.search(t)]
    candidates: set[int] = set()
    for text in filtered_texts:
        candidates |= _year_candidates_in_text(text)
    if len(candidates) == 1:
        return next(iter(candidates))
    return None  # ว่าง หรือกำกวม (หลายปีต่างกัน) → ปล่อย null + flag year_unknown


def _fiscal_year_texts(
    rows: list[tuple], region: TableRegion, sheet_name: str, rel_path: str
) -> list[str]:
    texts = [sheet_name, Path(rel_path).stem]
    start = max(0, region.header_row_idx - TITLE_LOOKBACK)
    for idx in range(start, region.header_row_idx + 1):
        for value in rows[idx]:
            if isinstance(value, str):
                texts.append(value)
    return texts


# -- row -> budget_line record ------------------------------------------------


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    return clean(str(value))


def _amount_to_baht(raw: object, unit_multiplier: int) -> int | None:
    if unit_multiplier == 1_000_000:
        return money.to_baht_from_million(raw)
    baht = money.parse_baht(raw)
    if baht is None:
        return None
    return baht * unit_multiplier


def _qty_to_float(raw: object) -> float | None:
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int | float):
        return float(raw)
    if isinstance(raw, str):
        baht = money.parse_baht(raw)  # ใช้ parser เดียวกัน (รองรับ comma/เลขไทย) ไม่ใช่หน่วยเงิน
        return float(baht) if baht is not None else None
    return None


def pyarrow_schema() -> pa.Schema:
    return pa.schema(
        [
            pa.field("source_id", pa.string()),
            pa.field("dataset", pa.string()),
            pa.field("fiscal_year_be", pa.int32()),
            pa.field("fiscal_year_ce", pa.int32()),
            pa.field("gov_level", pa.string()),
            pa.field("ministry", pa.string()),
            pa.field("agency", pa.string()),
            pa.field("province", pa.string()),
            pa.field("plan", pa.string()),
            pa.field("item_name_raw", pa.string()),
            pa.field("item_name", pa.string()),
            pa.field("item_qty", pa.float64()),
            pa.field("amount_thb", pa.int64()),
            pa.field("amount_unit_source", pa.string()),
            pa.field("source_path", pa.string()),
            pa.field("source_sheet", pa.string()),
            pa.field("source_row", pa.int32()),
            pa.field("source_doc_id", pa.string()),
            pa.field("quality_flags", pa.list_(pa.string())),
        ]
    )


@dataclass
class SheetMapResult:
    sheet_name: str
    mapped: bool
    reason: str | None = None
    n_rows: int = 0
    records: list[dict] = field(default_factory=list)
    subtotal_rows_skipped: int = 0


# แถวยอดรวม/สรุปยอด (ไม่ใช่รายการโครงการจริง) ที่มักปนอยู่ท้ายช่วงตารางที่ detect ได้
_SUBTOTAL_ITEM_NAME_PREFIXES: tuple[str, ...] = ("รวม", "ยอดรวม")


def _is_subtotal_row(item_name_raw: str, row: tuple, agency_col: int | None) -> bool:
    """แถว "รวม"/"รวม 2 แผน"/"ยอดรวม..." ปนอยู่ในช่วงตาราง — ไม่ใช่รายการโครงการจริง

    พบจริงใน กองทุนอนุรักษ์พลังงาน sheet `ปี 63` (แถว 361 "รวม 2 แผน"), `ปี 64 กลุ่ม 7`
    (แถว 989 "รวม"), `ปี 65 กลุ่ม 7` (แถว 991 "รวม") — ถ้านับเป็น budget_lines ปกติจะ
    double-count ยอดตอน aggregate ภายหลัง (เป็นผลรวมของแถวอื่นในตารางเดียวกัน)

    สัญญาณ: item_name ขึ้นต้นด้วย "รวม"/"ยอดรวม" **และ** คอลัมน์หน่วยงาน (ถ้ามี) ว่างเปล่า —
    แถวโครงการจริงมีหน่วยงานเสมอในไฟล์นี้ กันไม่ให้ตัดโครงการจริงที่ชื่อมีคำว่า "รวม" อยู่
    กลาง/ท้ายชื่อ (เช่น "...กลุ่มวิสาหกิจชุมชนรวมเกษตรยั่งยืน" — พบจริงในไฟล์เดียวกัน)
    """
    text = clean(item_name_raw)
    if not text.startswith(_SUBTOTAL_ITEM_NAME_PREFIXES):
        return False
    if agency_col is None:
        return True
    return _is_blank_cell(row[agency_col])


def _build_sheet_records(
    rows: list[tuple],
    region: TableRegion,
    mapping: dict[str, int],
    unit_label: str,
    unit_multiplier: int,
    *,
    rel_path: str,
    sheet_name: str,
    source_doc_id: str,
) -> tuple[list[dict], int]:
    """คืน `(records, subtotal_rows_skipped)`"""
    ministry_label = _detect_ministry_label(rows, region)
    fallback_agency = _agency_from_rel_path(rel_path)
    extra_text_for_gov_level = ministry_label or ""

    # ปีงบเป็นคุณสมบัติของทั้ง sheet (มาจากชื่อ sheet/title เดียวกันทุกแถว) — คำนวณครั้งเดียว
    fy_texts = _fiscal_year_texts(rows, region, sheet_name, rel_path)
    fiscal_year_be = _extract_fiscal_year_be(sheet_name, fy_texts)
    fiscal_year_ce = fiscal_year_be - 543 if fiscal_year_be is not None else None

    agency_col = mapping.get("agency")
    item_col = mapping.get("item_name")
    amount_col = mapping.get("amount_thb")
    plan_col = mapping.get("plan")
    province_col = mapping.get("province")
    qty_col = mapping.get("item_qty")

    records: list[dict] = []
    subtotal_rows_skipped = 0
    for row_idx in range(region.data_start_idx, region.data_end_idx):
        row = rows[row_idx]
        if all(_is_blank_cell(v) for v in row):
            continue

        item_name_raw = (
            nfc_only(str(row[item_col]))
            if item_col is not None and row[item_col] is not None
            else None
        )
        if item_name_raw is not None and clean(item_name_raw) in ("", "-"):
            item_name_raw = None
        amount_thb = (
            _amount_to_baht(row[amount_col], unit_multiplier) if amount_col is not None else None
        )

        if item_name_raw is None or amount_thb is None:
            # แถวไม่ครบ item+amount (เช่นแถวย่อย/หมายเหตุปนอยู่ในตาราง) — ข้ามเป็นรายแถว ไม่ล้มทั้ง sheet
            continue

        if _is_subtotal_row(item_name_raw, row, agency_col):
            subtotal_rows_skipped += 1
            continue

        flags: list[str] = []
        if fiscal_year_be is None:
            flags.append("year_unknown")

        # gov_level ต้องคำนวณ**ต่อแถว**ด้วย agency ที่ resolve แล้วจริง (ห้ามใช้ fallback_agency
        # ระดับ sheet ตรง ๆ — บั๊กจริงที่แก้ 19 ก.ย. 2569: เดิม gov_level ผูกกับ fallback_agency
        # ของทั้ง sheet ทำให้ 2,549/3,328 แถวที่เป็น อปท. จริง (เทศบาล/องค์การบริหารส่วน...)
        # ได้ gov_level="central" ผิดหมด)
        agency = _cell_to_str(row[agency_col]) or None if agency_col is not None else None
        agency = agency or fallback_agency
        gov_level = _gov_level_for_agency(agency, extra_text_for_gov_level)

        plan = _cell_to_str(row[plan_col]) or None if plan_col is not None else None
        province = _cell_to_str(row[province_col]) or None if province_col is not None else None
        item_qty = _qty_to_float(row[qty_col]) if qty_col is not None else None

        excel_row = row_idx + 1  # rows มาจาก iter_rows values_only เริ่ม index 0 = excel แถว 1
        records.append(
            {
                "source_id": compute_source_id(DATASET, rel_path, sheet_name, excel_row),
                "dataset": DATASET,
                "fiscal_year_be": fiscal_year_be,
                "fiscal_year_ce": fiscal_year_ce,
                "gov_level": gov_level,
                "ministry": ministry_label,
                "agency": agency,
                "province": province,
                "plan": plan,
                "item_name_raw": item_name_raw,
                "item_name": clean(item_name_raw),
                "item_qty": item_qty,
                "amount_thb": amount_thb,
                "amount_unit_source": unit_label,
                "source_path": rel_path,
                "source_sheet": sheet_name,
                "source_row": excel_row,
                "source_doc_id": source_doc_id,
                "quality_flags": flags,
            }
        )
    return records, subtotal_rows_skipped


def process_sheet(
    rows: list[tuple], sheet_name: str, *, rel_path: str, source_doc_id: str
) -> SheetMapResult:
    region = detect_table_region(rows)
    if region is None:
        return SheetMapResult(
            sheet_name=sheet_name, mapped=False, reason="ไม่พบตาราง (ไม่มี header+ข้อมูลตัวเลข)"
        )

    mapping = map_columns(region.header)
    missing = [f for f in REQUIRED_FIELDS if f not in mapping]
    if missing:
        return SheetMapResult(
            sheet_name=sheet_name, mapped=False, reason=f"map คอลัมน์ไม่ครบ (ขาด {missing})"
        )

    unit = detect_unit(rows, region)
    if unit is None:
        return SheetMapResult(
            sheet_name=sheet_name, mapped=False, reason="ไม่รู้หน่วยเงิน (ไม่พบ บาท/พันบาท/ล้านบาท)"
        )
    unit_label, unit_multiplier = unit

    records, subtotal_rows_skipped = _build_sheet_records(
        rows,
        region,
        mapping,
        unit_label,
        unit_multiplier,
        rel_path=rel_path,
        sheet_name=sheet_name,
        source_doc_id=source_doc_id,
    )
    if not records:
        reason = (
            "พบตาราง+หน่วยเงิน แต่มีแค่แถวยอดรวม ไม่มีรายการจริง"
            if subtotal_rows_skipped
            else "พบตาราง+หน่วยเงิน แต่ไม่มีแถวที่มีทั้ง item และ amount"
        )
        return SheetMapResult(
            sheet_name=sheet_name,
            mapped=False,
            reason=reason,
            subtotal_rows_skipped=subtotal_rows_skipped,
        )
    return SheetMapResult(
        sheet_name=sheet_name,
        mapped=True,
        n_rows=len(records),
        records=records,
        subtotal_rows_skipped=subtotal_rows_skipped,
    )


# -- unmapped sheet -> DocChunk atoms ----------------------------------------


def _non_blank_rows(rows: list[tuple]) -> list[tuple]:
    return [row for row in rows if not all(_is_blank_cell(v) for v in row)]


def _sheet_to_atoms(sheet_name: str, rows: list[tuple]) -> list[Atom]:
    atoms: list[Atom] = []
    table_key = f"sheet:{sheet_name}"
    for row in _non_blank_rows(rows):
        cells = [_cell_to_str(v) for v in row]
        atoms.append(Atom(line="\t".join(cells), table_row=cells, table_key=table_key))
    return atoms


# -- workbook I/O -------------------------------------------------------------


def _is_zip_magic(path: Path) -> bool:
    with open(path, "rb") as f:
        head = f.read(4)
    return head.startswith(b"PK")


def discover_committee_files(cfg: PipelineConfig) -> list[Path]:
    """เดินหาไฟล์ `.xlsx`/`.xls` ทั้งหมดใต้ `กมธ.ติดตามงบ` (เรียงตาม rel_path ให้ deterministic)"""
    root = cfg.raw_data_dir / COMMITTEE_ROOT
    if not root.is_dir():
        return []
    found: list[tuple[str, Path]] = []
    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in (".xlsx", ".xls"):
            rel_path = path.relative_to(cfg.raw_data_dir).as_posix()
            found.append((rel_path, path))
    found.sort(key=lambda pair: pair[0])
    return [abs_path for _rel, abs_path in found]


def _iter_workbook_sheets(path: Path) -> list[tuple[str, list[tuple]]]:
    """คืน `[(sheet_name, rows)]` — เลือก reader ตาม magic bytes จริง ไม่ใช่นามสกุลไฟล์"""
    if path.suffix.lower() == ".xlsx" or _is_zip_magic(path):
        return _iter_openpyxl_sheets(path)
    return _iter_xlrd_sheets(path)


def _iter_openpyxl_sheets(path: Path) -> list[tuple[str, list[tuple]]]:
    """เปิดด้วย **file object** (`rb`) เสมอ — ห้ามส่ง path string ตรง ๆ ให้ `load_workbook`

    ยืนยันจริง 19 ก.ย. 2569: `openpyxl.load_workbook(str(path), ...)` เช็คนามสกุลไฟล์ก่อนดู
    เนื้อไฟล์ → raise `"does not support the old .xls file format"` สำหรับ `องค์การโคนม/BI*.XLS`
    ทั้ง 13 ไฟล์ **แม้ magic bytes จะเป็น `PK` (zip/xlsx จริง)** เพราะ routing ใน
    `_iter_workbook_sheets` เลือกฟังก์ชันนี้ถูกแล้วแต่ตัวเรียก `load_workbook` เองยังเช็คนามสกุล
    จากอาร์กิวเมนต์ที่เป็น string — เปิดเป็น **binary file object** แทนจะข้าม path-extension
    check นี้ไปเลย (openpyxl ใช้ `zipfile` ตรวจเนื้อไฟล์แทน) ทดสอบแล้วเปิดได้ปกติทั้ง 13 ไฟล์
    """
    try:
        fh = open(path, "rb")  # noqa: SIM115 — ต้องคุมช่วงเปิดเองให้ยาวกว่า `with` บล็อกเดียว
    except OSError as exc:
        raise UnsupportedWorkbookError(str(exc)) from exc
    try:
        try:
            wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — openpyxl มี exception หลายชนิดเมื่อไฟล์เสีย
            raise UnsupportedWorkbookError(str(exc)) from exc
        try:
            sheets = []
            for name in wb.sheetnames:
                ws = wb[name]
                rows = list(ws.iter_rows(values_only=True))
                sheets.append((name, rows))
            return sheets
        finally:
            wb.close()
    finally:
        fh.close()


def _iter_xlrd_sheets(path: Path) -> list[tuple[str, list[tuple]]]:
    try:
        wb = xlrd.open_workbook(str(path))
    except Exception as exc:  # noqa: BLE001 — xlrd มี exception หลายชนิดเมื่อไฟล์เสีย/รูปแบบไม่รองรับ
        if shutil.which("soffice") is not None:
            raise UnsupportedWorkbookError(
                f"xlrd เปิดไม่ได้ ({exc}) — พบ soffice ใน PATH แต่ยังไม่ implement การแปลงผ่าน LibreOffice"
            ) from exc
        raise UnsupportedWorkbookError(
            f"xlrd เปิด .xls ไม่ได้ ({exc}) — เครื่อง dev ไม่มี LibreOffice ให้แปลง"
        ) from exc
    sheets = []
    for name in wb.sheet_names():
        ws = wb.sheet_by_name(name)
        rows = [tuple(ws.row_values(r)) for r in range(ws.nrows)]
        sheets.append((name, rows))
    return sheets


# -- orchestration ------------------------------------------------------


@dataclass
class CommitteeFileResult:
    rel_path: str
    doc_id: str
    mapped_rows: int = 0
    sheets_mapped: list[str] = field(default_factory=list)
    sheets_unmapped: list[tuple[str, str]] = field(default_factory=list)  # (sheet, reason)
    subtotal_rows_skipped: int = 0
    parquet_path: Path | None = None
    doc_chunk_path: Path | None = None
    skipped: bool = False
    skip_reason: str | None = None


@dataclass
class CommitteeExtractSummary:
    results: list[CommitteeFileResult] = field(default_factory=list)

    @property
    def total_mapped_rows(self) -> int:
        return sum(r.mapped_rows for r in self.results)


def _parquet_cache_path(cfg: PipelineConfig, doc_id: str, cache_dir: Path | None) -> Path:
    base = cache_dir if cache_dir is not None else cfg.cache_dir / COMMITTEE_CACHE_DIRNAME
    return base / f"{doc_id}.parquet"


def _docs_cache_path(cfg: PipelineConfig, doc_id: str, docs_cache_dir: Path | None) -> Path:
    base = docs_cache_dir if docs_cache_dir is not None else cfg.cache_dir / "docs"
    return base / f"{doc_id}.json.gz"


def extract_committee_file(
    cfg: PipelineConfig,
    path: Path,
    *,
    cache_dir: Path | None = None,
    docs_cache_dir: Path | None = None,
) -> CommitteeFileResult:
    rel_path = path.relative_to(cfg.raw_data_dir).as_posix()
    doc_id = doc_id_for_path(rel_path)
    result = CommitteeFileResult(rel_path=rel_path, doc_id=doc_id)

    try:
        sheets = _iter_workbook_sheets(path)
    except UnsupportedWorkbookError as exc:
        result.skipped = True
        result.skip_reason = str(exc)
        return result

    all_records: list[dict] = []
    doc_chunks: list[DocChunk] = []
    chunk_no = 0

    for sheet_name, rows in sheets:
        if not rows:
            continue
        sheet_result = process_sheet(rows, sheet_name, rel_path=rel_path, source_doc_id=doc_id)
        result.subtotal_rows_skipped += sheet_result.subtotal_rows_skipped
        if sheet_result.mapped:
            result.sheets_mapped.append(sheet_name)
            all_records.extend(sheet_result.records)
        else:
            result.sheets_unmapped.append((sheet_name, sheet_result.reason or "ไม่ทราบสาเหตุ"))
            atoms = _sheet_to_atoms(sheet_name, rows)
            for text, tables in chunk_atoms(atoms):
                # แนบชื่อ sheet ให้ทุกตารางในก้อนนี้ (atoms ของ sheet เดียวกันเสมอ — ดู `_sheet_to_atoms`)
                for table in tables:
                    table["sheet"] = sheet_name
                doc_chunks.append(
                    {
                        "doc_id": doc_id,
                        "page": None,
                        "chunk_no": chunk_no,
                        "text": text,
                        "tables": tables,
                    }
                )
                chunk_no += 1

    result.mapped_rows = len(all_records)

    if all_records:
        out_path = cfg.assert_writable_path(_parquet_cache_path(cfg, doc_id, cache_dir))
        out_path.parent.mkdir(parents=True, exist_ok=True)
        table = pa.Table.from_pylist(all_records, schema=pyarrow_schema())
        pq.write_table(table, str(out_path), compression="zstd")
        result.parquet_path = out_path

    if doc_chunks:
        docs_path = cfg.assert_writable_path(_docs_cache_path(cfg, doc_id, docs_cache_dir))
        write_doc_chunks_gz(docs_path, doc_chunks)
        result.doc_chunk_path = docs_path

    return result


def extract_committee_xlsx(
    cfg: PipelineConfig, *, cache_dir: Path | None = None, docs_cache_dir: Path | None = None
) -> CommitteeExtractSummary:
    """extract xlsx/xls ของ กมธ.ติดตามงบ ทั้งหมด (03 §4.4) → `.cache/committee/` + `.cache/docs/`"""
    results = [
        extract_committee_file(cfg, path, cache_dir=cache_dir, docs_cache_dir=docs_cache_dir)
        for path in discover_committee_files(cfg)
    ]
    return CommitteeExtractSummary(results=results)
