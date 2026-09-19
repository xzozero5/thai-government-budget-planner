"""T-108: docx/pptx → `DocChunk` (03-DATA-PIPELINE.md §3.3, §4.5)

`DocChunk` ตาม 03 §3.3 เป๊ะ ๆ: `{doc_id, page, chunk_no, text, tables}`
- pptx: `page` = ลำดับ slide (1-based); เนื้อความของ 1 slide (shapes + notes) chunk ต่อเนื่อง
  ~800-1,200 ตัวอักษร ไม่ตัดกลาง "หน่วยข้อความอะตอม" (บรรทัด/แถวตาราง) — ถ้า slide เดียวมีหลาย
  chunk ทุก chunk มี `page` เดิม, `chunk_no` เพิ่มขึ้น
- docx: ไม่มีเลขหน้าจริง (python-docx ไม่ทราบ pagination) → `page: null` เสมอ, ใช้ `chunk_no`
  อ่านตามลำดับจริงในเอกสาร (paragraph/table สลับกันได้ ตาม XML body order) — คงเนื้อหาตาราง
  เป็นแถว ๆ ไม่ตัดกลางแถว (`tables: [{"rows": [[...]]}]`)

โมดูลนี้ยังเป็นที่มาของ `DocChunk`/`chunk_atoms`/`write_doc_chunks_gz` ที่ `extract/committee_xlsx.py`
เรียกใช้ร่วม (ตาราง xlsx ที่ map ไม่ได้ก็เก็บเป็น `DocChunk` แบบเดียวกัน) — ห้ามสร้างซ้ำที่อื่น (N7)
T-109 (`extract/pdf_text.py`) ใช้ `DocChunk`/`write_doc_chunks_gz` เดียวกันนี้เช่นกัน (main thread
จะรวม `tgbp extract --dataset office|pdf|committee` ภายหลัง)
"""

from __future__ import annotations

import gzip
import json
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import NotRequired, TypedDict

import docx
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from pptx import Presentation

from tgbp_pipeline.config import PipelineConfig
from tgbp_pipeline.normalize.thai_text import clean, nfc_only
from tgbp_pipeline.util.hash import doc_id_for_path

DOCS_CACHE_DIRNAME = "docs"
CHUNK_CHAR_LIMIT = 1200

_OFFICE_EXTENSIONS: tuple[str, ...] = (".docx", ".pptx")


class DocTable(TypedDict):
    """หนึ่งตารางที่แนบกับ `DocChunk` — `sheet`/`page` ใส่มาก็ต่อเมื่อมีแนวคิดนั้นจริง"""

    sheet: NotRequired[str]
    page: NotRequired[int]
    rows: list[list[str]]


class DocChunk(TypedDict):
    """ตาม 03-DATA-PIPELINE.md §3.3 เป๊ะ ๆ — ห้ามเพิ่ม field เกินสเปกที่ main thread ไม่รู้จัก"""

    doc_id: str
    page: int | None
    chunk_no: int
    text: str
    tables: list[DocTable]


@dataclass
class Atom:
    """หน่วยข้อความอะตอมหนึ่งชิ้น (บรรทัด/แถวตาราง) ที่ห้ามถูกตัดครึ่งระหว่าง chunk"""

    line: str
    table_row: list[str] | None = None
    table_key: object | None = None  # object เดียวกัน (เช่น id ของตาราง) = แถวเดียวกันของตาราง


def chunk_atoms(
    atoms: list[Atom], limit: int = CHUNK_CHAR_LIMIT
) -> list[tuple[str, list[DocTable]]]:
    """รวม atom ต่อเนื่องให้ได้ก้อนข้อความ ~`limit` ตัวอักษร โดยไม่ตัดกลาง atom เดียว

    คืนลิสต์ของ (text, tables) — `tables` รวมแถวของตารางเดียวกันที่อยู่ติดกันเป็นรายการเดียว
    """
    chunks: list[tuple[str, list[DocTable]]] = []
    current_lines: list[str] = []
    current_len = 0
    current_tables: list[DocTable] = []
    current_table_key: object | None = None

    def flush() -> None:
        nonlocal current_lines, current_len, current_tables, current_table_key
        if not current_lines:
            return
        chunks.append(("\n".join(current_lines), current_tables))
        current_lines = []
        current_len = 0
        current_tables = []
        current_table_key = None

    for atom in atoms:
        separator_cost = 1 if current_lines else 0
        projected_len = current_len + separator_cost + len(atom.line)
        if current_lines and projected_len > limit:
            flush()
            projected_len = len(atom.line)
        current_lines.append(atom.line)
        current_len = projected_len
        if atom.table_row is not None:
            if current_table_key == atom.table_key and current_tables:
                current_tables[-1]["rows"].append(atom.table_row)
            else:
                current_tables.append({"rows": [atom.table_row]})
                current_table_key = atom.table_key
    flush()
    return chunks


def write_doc_chunks_gz(path: Path, chunks: list[DocChunk]) -> None:
    """เขียน `docs/{doc_id}.json.gz` แบบ deterministic (mtime=0, ไม่มี FNAME) — รันซ้ำได้ byte เท่ากัน"""
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(chunks, ensure_ascii=False).encode("utf-8")
    with open(path, "wb") as raw:
        with gzip.GzipFile(filename="", fileobj=raw, mode="wb", mtime=0) as gz:
            gz.write(payload)


def read_doc_chunks_gz(path: Path) -> list[DocChunk]:
    with gzip.open(path, "rb") as gz:
        return json.loads(gz.read().decode("utf-8"))


def doc_chunks_cache_path(cfg: PipelineConfig, doc_id: str, cache_dir: Path | None = None) -> Path:
    base = cache_dir if cache_dir is not None else cfg.cache_dir / DOCS_CACHE_DIRNAME
    return base / f"{doc_id}.json.gz"


def discover_office_files(cfg: PipelineConfig) -> list[Path]:
    """เดินหาไฟล์ `.docx`/`.pptx` ทั้งหมดใต้ raw_data_dir (เรียงตาม rel_path ให้ deterministic)"""
    found: list[tuple[str, Path]] = []
    for dirpath, _dirnames, filenames in os.walk(cfg.raw_data_dir):
        for name in filenames:
            if Path(name).suffix.lower() in _OFFICE_EXTENSIONS:
                abs_path = Path(dirpath) / name
                rel_path = abs_path.relative_to(cfg.raw_data_dir).as_posix()
                found.append((rel_path, abs_path))
    found.sort(key=lambda pair: pair[0])
    return [abs_path for _rel, abs_path in found]


def _cell_text(text: str) -> str:
    return clean(text).replace("\t", " ").replace("\n", " ")


# -- docx -----------------------------------------------------------------


def _iter_docx_block_items(document: docx.document.Document):
    """iterate paragraph/table ตามลำดับจริงใน body XML (recipe มาตรฐานของ python-docx)"""
    parent_elm = document.element.body
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, document)


def _docx_table_to_atoms(table: DocxTable) -> list[Atom]:
    table_key = id(table)
    atoms: list[Atom] = []
    for row in table.rows:
        cells = [_cell_text(cell.text) for cell in row.cells]
        line = "\t".join(cells)
        if not any(cells):
            continue
        atoms.append(Atom(line=line, table_row=cells, table_key=table_key))
    return atoms


def _docx_to_atoms(document: docx.document.Document) -> list[Atom]:
    atoms: list[Atom] = []
    for block in _iter_docx_block_items(document):
        if isinstance(block, DocxParagraph):
            text = nfc_only(block.text).strip()
            if text:
                atoms.append(Atom(line=text))
        elif isinstance(block, DocxTable):
            atoms.extend(_docx_table_to_atoms(block))
    return atoms


def extract_docx_file(doc_id: str, path: Path) -> list[DocChunk]:
    """หนึ่งไฟล์ `.docx` → `DocChunk[]` (`page: null` เสมอ — python-docx ไม่รู้เลขหน้า)"""
    document = docx.Document(str(path))
    atoms = _docx_to_atoms(document)
    chunks: list[DocChunk] = []
    for chunk_no, (text, tables) in enumerate(chunk_atoms(atoms)):
        chunks.append(
            {"doc_id": doc_id, "page": None, "chunk_no": chunk_no, "text": text, "tables": tables}
        )
    return chunks


# -- pptx -------------------------------------------------------------------


def _pptx_table_atoms(table) -> list[Atom]:
    table_key = id(table)
    atoms: list[Atom] = []
    for row in table.rows:
        cells = [_cell_text(cell.text) for cell in row.cells]
        if not any(cells):
            continue
        atoms.append(Atom(line="\t".join(cells), table_row=cells, table_key=table_key))
    return atoms


def _slide_atoms(slide) -> list[Atom]:
    atoms: list[Atom] = []
    for shape in slide.shapes:
        if getattr(shape, "has_table", False) and shape.has_table:
            atoms.extend(_pptx_table_atoms(shape.table))
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = nfc_only(shape.text_frame.text).strip()
            if text:
                for line in text.splitlines():
                    stripped = clean(line)
                    if stripped:
                        atoms.append(Atom(line=stripped))
    if slide.has_notes_slide:
        notes_frame = slide.notes_slide.notes_text_frame
        if notes_frame is not None:
            notes_text = nfc_only(notes_frame.text).strip()
            if notes_text:
                atoms.append(Atom(line=f"[notes] {clean(notes_text)}"))
    return atoms


def extract_pptx_file(doc_id: str, path: Path) -> list[DocChunk]:
    """หนึ่งไฟล์ `.pptx` → `DocChunk[]` (`page` = ลำดับ slide เริ่มที่ 1)"""
    presentation = Presentation(str(path))
    chunks: list[DocChunk] = []
    for slide_no, slide in enumerate(presentation.slides, start=1):
        atoms = _slide_atoms(slide)
        if not atoms:
            continue
        for chunk_no, (text, tables) in enumerate(chunk_atoms(atoms)):
            chunks.append(
                {
                    "doc_id": doc_id,
                    "page": slide_no,
                    "chunk_no": chunk_no,
                    "text": text,
                    "tables": tables,
                }
            )
    return chunks


# -- orchestration ------------------------------------------------------


@dataclass
class OfficeTextResult:
    rel_path: str
    doc_id: str
    kind: str
    n_chunks: int
    cache_path: Path
    error: str | None = None


@dataclass
class OfficeTextSummary:
    results: list[OfficeTextResult] = field(default_factory=list)


def extract_office_file(
    cfg: PipelineConfig, path: Path, cache_dir: Path | None = None
) -> OfficeTextResult:
    rel_path = path.relative_to(cfg.raw_data_dir).as_posix()
    doc_id = doc_id_for_path(rel_path)
    kind = path.suffix.lower().lstrip(".")
    out_path = cfg.assert_writable_path(doc_chunks_cache_path(cfg, doc_id, cache_dir))

    try:
        if kind == "docx":
            chunks = extract_docx_file(doc_id, path)
        elif kind == "pptx":
            chunks = extract_pptx_file(doc_id, path)
        else:  # pragma: no cover — discover_office_files กรองแค่ 2 นามสกุลนี้
            raise ValueError(f"ไม่รองรับนามสกุล {kind!r}")
    except Exception as exc:  # noqa: BLE001 — ไฟล์ office เสีย/รูปแบบแปลกมีหลายชนิด exception
        return OfficeTextResult(
            rel_path=rel_path,
            doc_id=doc_id,
            kind=kind,
            n_chunks=0,
            cache_path=out_path,
            error=str(exc),
        )

    write_doc_chunks_gz(out_path, chunks)
    return OfficeTextResult(
        rel_path=rel_path, doc_id=doc_id, kind=kind, n_chunks=len(chunks), cache_path=out_path
    )


def extract_office_text(cfg: PipelineConfig, *, cache_dir: Path | None = None) -> OfficeTextSummary:
    """extract docx/pptx ทั้งหมดใต้ raw_data_dir → `.cache/docs/{doc_id}.json.gz` (03 §4.5)"""
    results = [
        extract_office_file(cfg, path, cache_dir=cache_dir) for path in discover_office_files(cfg)
    ]
    return OfficeTextSummary(results=results)
