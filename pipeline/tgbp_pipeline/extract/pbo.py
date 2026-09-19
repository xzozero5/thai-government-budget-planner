"""T-105: `PBO/25xx.xlsx` — ผลเบิกจ่ายรายรายการ ปีงบ 2558-2568 (03 §4.1, 02 §A1)

Stream ทีละ `CHUNK_ROWS` แถวด้วย `openpyxl.load_workbook(read_only=True)` แล้วเขียน
`pipeline/.cache/pbo/{year}.parquet` (zstd) — **ยังไม่ทำ item_parser/org_master** (T-102..T-104
ทำแยก) จึงมีแค่ฟิลด์ที่รู้ได้ตรง ๆ จากไฟล์ (ministry/agency เป็นชื่อดิบ ยังไม่ map เป็นรหัส)

Gotchas ที่ยืนยันแล้ว (02 §A1) ที่โมดูลนี้ต้องจัดการ:
- เลือก sheet ด้วย prefix `เบิกจ่ายภาพรวมทุกมิติ` (2561/2568 มี `Sheet1` pivot นำหน้า)
- ตัดที่ 22 คอลัมน์แรกเสมอ (`iter_rows(max_col=22, ...)`) — ปี 2567 มี max_column 16384
- แถว `Grand Total` ตรวจด้วย**เนื้อหา** ไม่ใช่เลขแถว (2561/2567/2568 ไม่มีแถวนี้)
- ค่าว่าง = สตริง `'-'`; เงินหน่วยล้านบาท → บาท int ผ่าน `normalize.money.to_baht_from_million`
- oracle (สำหรับ `tgbp_pipeline.validate.check_v1`) เก็บเป็น**ล้านบาทตามไฟล์** ใน `oracle.json`:
    - มี Grand Total แถว 2 → kind `grand_total`
    - ไม่มี Grand Total แต่มี `Sheet1` แคบ (2 คอลัมน์: ป้ายชื่อแถว/ผลรวม) → kind `sheet1_ministry`
      (2561 — pivot รายกระทรวงทั้งหมด)
    - ไม่มี Grand Total แต่มี `Sheet1` กว้าง (header 2 ชั้น ประจำ/ลงทุน × 3 คอลัมน์) →
      kind `sheet1_total` (2568 — pivot กรองเฉพาะสำนักนายกรัฐมนตรี เท่านั้น)
    - ไม่มีทั้งคู่ → kind `none` (2567 — ไม่มี oracle ในไฟล์เลย)

แก้ 19 ก.ย. 2569 (รอบ 2 — หลัง main thread ตรวจ cache จริง):
- **`--limit-rows`** เขียนไปไฟล์ `{year}.partial.parquet` เสมอ (ไม่ทับ `{year}.parquet` เต็ม)
  และ**ไม่แตะ** `oracle.json`/`headers.json` เลย (กัน rawdata test ที่จำกัดแถวทับ cache จริง
  ที่เคยเกิดขึ้นจริงกับ `2566.parquet` — เหลือแค่ 5,000 แถวจนกว่าจะ extract ใหม่)
- **แถวเสียจากคอลัมน์เลื่อน** (พบจริงใน `PBO/2562.xlsx` แถว 226-227 — ดู
  `pipeline/data/known_source_gaps.yaml`): ถ้าคอลัมน์เงินคอลัมน์ใดเป็นข้อความที่ไม่ใช่ตัวเลข/`'-'`
  หรือ `abs(บาท) > CORRUPT_ROW_ABS_BAHT_THRESHOLD` (1e13 บาท ≈ 3 เท่าของงบประเทศทั้งปี) →
  ตั้ง**ทุก**ฟิลด์เงินของแถวนั้นเป็น `null` + flag `corrupt_row` (คง `item_name_raw`/มิติอื่นไว้
  ตามจริงเพื่อ citation) — `check_v1` จึงไม่นับแถวนี้โดยอัตโนมัติ (sum ข้าม null)
"""

from __future__ import annotations

import json
import time
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from tgbp_pipeline.config import PipelineConfig
from tgbp_pipeline.normalize import money
from tgbp_pipeline.normalize.thai_text import clean, nfc_only
from tgbp_pipeline.util.hash import doc_id_for_path
from tgbp_pipeline.util.hash import source_id as compute_source_id

DATASET = "pbo_disbursement"
SHEET_NAME_PREFIX = "เบิกจ่ายภาพรวมทุกมิติ"
MAX_COLS = 22
CHUNK_ROWS = 50_000
AMOUNT_UNIT_SOURCE = "million_thb"

# 02 §A1 — ปี PBO ที่มีให้จริง (2558..2568)
ALL_YEARS: tuple[int, ...] = tuple(range(2558, 2569))

# ปี 2567 ไม่มี oracle ในไฟล์เลย (02 §A1) — ตรวจจำนวนแถวเทียบค่านี้แบบ soft แทน
YEAR_NO_ORACLE_EXPECTED_ROWS: dict[int, int] = {2567: 221_571}

# header 22 คอลัมน์แรก — ยืนยันเหมือนกันทุกปี (02 §A1, สแกนจริง 19 ก.ย. 2569)
EXPECTED_HEADER: tuple[str, ...] = (
    "ปีงบประมาณ",
    "กระทรวง",
    "หน่วยงาน",
    "ยุทธศาสตร์การจัดสรร",
    "แผนงาน",
    "ผลผลิต/โครงการ",
    "งาน/โครงการ",
    "งบรายจ่าย",
    "รายจ่ายประจำ/ลงทุน",
    "ชื่อรหัสงบประมาณ",
    "พรบ. (ล้านบาท)",
    "งบฯ หลังโอน/ปป. ทั้งสิ้น (ล้านบาท)",
    "PO ทั้งสิ้น (ล้านบาท)",
    "เบิกจ่ายทั้งสิ้น (ล้านบาท)",
    "เบิกจ่ายรวม PO (ล้านบาท)",
    "เงินกันฯ สุทธิ (ล้านบาท)",
    "เบิกจ่ายเหลื่อมปี (ล้านบาท)",
    "คงเหลือกรณีมีหนี้ผูกพัน (ล้านบาท)",
    "คงเหลือกรณีไม่มีหนี้ผูกพัน (ล้านบาท)",
    "คงเหลือ สรก.อยู่ระหว่างดำเนินการ (ล้านบาท)",
    "คงเหลืออยู่ระหว่างกันและขยายรวม (ล้านบาท)",
    "คงเหลือรวม (ล้านบาท)",
)

# (0-based index ใน row 22 คอลัมน์, ชื่อฟิลด์ cache, header ไทยตรง — ใช้ทั้งอ่าน 22-col sheet
# และอ่านแถว Grand Total/Sheet1 total ที่มีคอลัมน์เดียวกัน)
MONEY_COLUMNS: tuple[tuple[int, str, str], ...] = (
    (10, "amount_thb", "พรบ. (ล้านบาท)"),
    (11, "revised_thb", "งบฯ หลังโอน/ปป. ทั้งสิ้น (ล้านบาท)"),
    (12, "po_thb", "PO ทั้งสิ้น (ล้านบาท)"),
    (13, "disbursed_thb", "เบิกจ่ายทั้งสิ้น (ล้านบาท)"),
    (14, "disbursed_incl_po_thb", "เบิกจ่ายรวม PO (ล้านบาท)"),
    (15, "reserved_thb", "เงินกันฯ สุทธิ (ล้านบาท)"),
    (16, "carryover_thb", "เบิกจ่ายเหลื่อมปี (ล้านบาท)"),
    (17, "remaining_committed_thb", "คงเหลือกรณีมีหนี้ผูกพัน (ล้านบาท)"),
    (18, "remaining_uncommitted_thb", "คงเหลือกรณีไม่มีหนี้ผูกพัน (ล้านบาท)"),
    (19, "remaining_in_progress_thb", "คงเหลือ สรก.อยู่ระหว่างดำเนินการ (ล้านบาท)"),
    (20, "remaining_reserved_extended_thb", "คงเหลืออยู่ระหว่างกันและขยายรวม (ล้านบาท)"),
    (21, "remaining_total_thb", "คงเหลือรวม (ล้านบาท)"),
)

_CAPITAL_TYPE_MAP: dict[str, bool] = {
    "รายจ่ายลงทุน": True,
    "รายจ่ายประจำ": False,
}

ORACLE_FILENAME = "oracle.json"
HEADERS_FILENAME = "headers.json"
PARTIAL_SUFFIX = ".partial.parquet"

# แถว "คอลัมน์เลื่อน" จริงพบใน PBO/2562.xlsx แถว 226 — พรบ. = 607,006,457,005 ล้านบาท
# (มากกว่างบประเทศทั้งปีหลายแสนเท่า) — 1e13 บาท ≈ 3 เท่าของงบประเทศทั้งปี (~3-3.3 ล้านล้านบาท)
CORRUPT_ROW_ABS_BAHT_THRESHOLD = 10_000_000_000_000  # 1e13 บาท = 10 ล้านล้านบาท


class PboHeaderMismatchError(ValueError):
    """header 22 คอลัมน์แรกของไฟล์ไม่ตรงกับปีอ้างอิง (2558) — 03 §4.1"""


class PboSheetNotFoundError(ValueError):
    """ไม่พบ sheet ที่ขึ้นต้นด้วย `เบิกจ่ายภาพรวมทุกมิติ` (หรือเจอมากกว่า 1 sheet)"""


@dataclass
class PboYearResult:
    year: int
    file_path: Path
    rel_path: str
    sheet_name: str
    rows_written: int
    n_grand_total_rows: int
    flag_counts: Counter
    cache_path: Path
    cache_bytes: int
    elapsed_seconds: float
    oracle: dict


@dataclass
class _RowFields:
    fiscal_year_be: int | None
    fiscal_year_ce: int | None
    ministry: str | None
    agency: str | None
    strategy: str | None
    plan: str | None
    output_project: str | None
    activity: str | None
    budget_type: str | None
    is_capital: bool | None
    capital_type_raw: str | None
    item_name_raw: str | None
    money_values: dict[str, int | None]
    flags: list[str] = field(default_factory=list)


def discover_year_files(cfg: PipelineConfig, years: list[int] | None = None) -> dict[int, Path]:
    """หาไฟล์ `PBO/{year}.xlsx` ที่มีอยู่จริงใต้ `raw_data_dir` (เรียงตามปี)"""
    wanted = years if years is not None else list(ALL_YEARS)
    pbo_dir = cfg.raw_data_dir / "PBO"
    result: dict[int, Path] = {}
    for year in wanted:
        path = pbo_dir / f"{year}.xlsx"
        if path.is_file():
            result[year] = path
    return dict(sorted(result.items()))


def _pick_data_sheet(sheet_names: list[str], year: int, path: Path) -> str:
    candidates = [name for name in sheet_names if name.startswith(SHEET_NAME_PREFIX)]
    if not candidates:
        raise PboSheetNotFoundError(
            f"ไม่พบ sheet ที่ขึ้นต้นด้วย '{SHEET_NAME_PREFIX}' ใน {path} (ปี {year}); "
            f"sheet จริง: {sheet_names}"
        )
    if len(candidates) > 1:
        raise PboSheetNotFoundError(
            f"เจอ sheet ที่ขึ้นต้นด้วย '{SHEET_NAME_PREFIX}' มากกว่า 1 sheet ใน {path} "
            f"(ปี {year}): {candidates}"
        )
    return candidates[0]


def _validate_header(header: tuple, year: int, path: Path) -> None:
    trimmed = tuple(header[: len(EXPECTED_HEADER)])
    if trimmed == EXPECTED_HEADER:
        return
    mismatches = [
        f"คอลัมน์ {i}: คาด {expected!r} ได้ {actual!r}"
        for i, (expected, actual) in enumerate(zip(EXPECTED_HEADER, trimmed, strict=False), start=1)
        if expected != actual
    ]
    if len(trimmed) < len(EXPECTED_HEADER):
        mismatches.append(f"มีแค่ {len(trimmed)} คอลัมน์ (ต้องการ {len(EXPECTED_HEADER)})")
    raise PboHeaderMismatchError(
        f"header 22 คอลัมน์แรกของ PBO ปี {year} ({path}) ไม่ตรงกับปีอ้างอิง 2558: " + "; ".join(mismatches)
    )


def _row_is_grand_total(row: tuple) -> bool:
    """ตรวจแถว Grand Total ด้วยเนื้อหา (03 §4.1) — คอลัมน์มิติ 1-9 (index 0-8) เป็น 'Grand Total'"""
    return any(isinstance(v, str) and v.strip() == "Grand Total" for v in row[:9])


def _text_or_none(value: object) -> str | None:
    if value is None:
        return None
    text = clean(str(value))
    if text in ("", "-"):
        return None
    return text


def _item_name_raw_and_empty_flag(value: object) -> tuple[str | None, bool]:
    """`item_name_raw` — NFC เท่านั้น (ห้าม transliterate/ตัดคำอื่น ๆ ของต้นฉบับ)"""
    if value is None:
        return None, True
    text = nfc_only(str(value))
    if text.strip() in ("", "-"):
        return None, True
    return text, False


def _oracle_value(value: object) -> float | None:
    """ค่าตัวเลขจาก Grand Total/Sheet1 (ล้านบาทตามไฟล์) — `'-'`/ว่าง → None"""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean(str(value)).strip()
    if text in ("", "-"):
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None


def _is_unparseable_money_text(raw: object) -> bool:
    """`True` ถ้า raw เป็นข้อความที่ไม่ใช่ทั้งตัวเลขและไม่ใช่ `'-'`/ว่าง (เช่น 'งบลงทุน' หลุดมา
    ในคอลัมน์เงิน — พบจริงใน PBO/2562.xlsx แถว 226 คอลัมน์ 'PO ทั้งสิ้น'/'เบิกจ่ายทั้งสิ้น')
    """
    if raw is None or isinstance(raw, (int, float)):
        return False
    if isinstance(raw, bool):
        return True
    text = clean(str(raw)).strip()
    if text in ("", "-"):
        return False
    try:
        float(text.replace(",", ""))
    except ValueError:
        return True
    return False


def _row_has_corrupt_money_cell(row: tuple) -> bool:
    """ตรวจว่าแถวนี้มีคอลัมน์เงินที่ 'เลื่อนคอลัมน์'/เสีย — ข้อความปนตัวเลข หรือค่าใหญ่เกินจริง"""
    for idx, _field_name, _label in MONEY_COLUMNS:
        raw = row[idx]
        if _is_unparseable_money_text(raw):
            return True
        baht = money.to_baht_from_million(raw)
        if baht is not None and abs(baht) > CORRUPT_ROW_ABS_BAHT_THRESHOLD:
            return True
    return False


def _build_row_fields(row: tuple, year: int) -> _RowFields:
    flags: list[str] = []

    fiscal_year_be: int | None
    try:
        fiscal_year_be = int(row[0])
    except (TypeError, ValueError):
        fiscal_year_be = None
    if fiscal_year_be != year:
        flags.append("year_mismatch")
    fiscal_year_ce = fiscal_year_be - 543 if fiscal_year_be is not None else None

    capital_type_raw = _text_or_none(row[8])
    is_capital = _CAPITAL_TYPE_MAP.get(capital_type_raw) if capital_type_raw else None

    item_name_raw, empty_name = _item_name_raw_and_empty_flag(row[9])
    if empty_name:
        flags.append("empty_item_name")

    is_corrupt = _row_has_corrupt_money_cell(row)
    if is_corrupt:
        # คอลัมน์เงินเลื่อน/เสีย (03 §4.1 แก้ 19 ก.ย. 2569 รอบ 2) — null ทุกฟิลด์เงิน
        # ของแถวนี้ทิ้งทั้งหมด ไม่พยายามเดาว่าค่าไหน "ถูก" (คงมิติ/item_name_raw ไว้เพื่อ citation)
        money_values = {field_name: None for _idx, field_name, _label in MONEY_COLUMNS}
        flags.append("corrupt_row")
    else:
        money_values = {}
        has_negative = False
        for idx, field_name, _label in MONEY_COLUMNS:
            baht = money.to_baht_from_million(row[idx])
            money_values[field_name] = baht
            if baht is not None and baht < 0:
                has_negative = True
        if has_negative:
            flags.append("negative_amount")

    return _RowFields(
        fiscal_year_be=fiscal_year_be,
        fiscal_year_ce=fiscal_year_ce,
        ministry=_text_or_none(row[1]),
        agency=_text_or_none(row[2]),
        strategy=_text_or_none(row[3]),
        plan=_text_or_none(row[4]),
        output_project=_text_or_none(row[5]),
        activity=_text_or_none(row[6]),
        budget_type=_text_or_none(row[7]),
        is_capital=is_capital,
        capital_type_raw=capital_type_raw,
        item_name_raw=item_name_raw,
        money_values=money_values,
        flags=flags,
    )


def _build_record(
    row: tuple,
    year: int,
    excel_row: int,
    rel_path: str,
    sheet_name: str,
    source_doc_id: str,
) -> tuple[dict, list[str]]:
    fields = _build_row_fields(row, year)

    revised = fields.money_values.get("revised_thb")
    disbursed = fields.money_values.get("disbursed_thb")
    disbursement_rate = (disbursed / revised) if revised and disbursed is not None else None

    record = {
        "source_id": compute_source_id(DATASET, rel_path, sheet_name, excel_row),
        "dataset": DATASET,
        "fiscal_year_be": fields.fiscal_year_be,
        "fiscal_year_ce": fields.fiscal_year_ce,
        "ministry": fields.ministry,
        "agency": fields.agency,
        "strategy": fields.strategy,
        "plan": fields.plan,
        "output_project": fields.output_project,
        "activity": fields.activity,
        "budget_type": fields.budget_type,
        "is_capital": fields.is_capital,
        "capital_type_raw": fields.capital_type_raw,
        "item_name_raw": fields.item_name_raw,
        **fields.money_values,
        "disbursement_rate": disbursement_rate,
        "amount_unit_source": AMOUNT_UNIT_SOURCE,
        "source_path": rel_path,
        "source_sheet": sheet_name,
        "source_row": excel_row,
        "source_doc_id": source_doc_id,
        "quality_flags": fields.flags,
    }
    return record, fields.flags


def pyarrow_schema() -> pa.Schema:
    money_fields = [pa.field(name, pa.int64()) for _, name, _ in MONEY_COLUMNS]
    return pa.schema(
        [
            pa.field("source_id", pa.string()),
            pa.field("dataset", pa.string()),
            pa.field("fiscal_year_be", pa.int32()),
            pa.field("fiscal_year_ce", pa.int32()),
            pa.field("ministry", pa.string()),
            pa.field("agency", pa.string()),
            pa.field("strategy", pa.string()),
            pa.field("plan", pa.string()),
            pa.field("output_project", pa.string()),
            pa.field("activity", pa.string()),
            pa.field("budget_type", pa.string()),
            pa.field("is_capital", pa.bool_()),
            pa.field("capital_type_raw", pa.string()),
            pa.field("item_name_raw", pa.string()),
            *money_fields,
            pa.field("disbursement_rate", pa.float64()),
            pa.field("amount_unit_source", pa.string()),
            pa.field("source_path", pa.string()),
            pa.field("source_sheet", pa.string()),
            pa.field("source_row", pa.int32()),
            pa.field("source_doc_id", pa.string()),
            pa.field("quality_flags", pa.list_(pa.string())),
        ]
    )


def _write_chunk(
    buffer: list[dict], cache_path: Path, writer: pq.ParquetWriter | None
) -> pq.ParquetWriter:
    schema = pyarrow_schema()
    table = pa.Table.from_pylist(buffer, schema=schema)
    if writer is None:
        writer = pq.ParquetWriter(str(cache_path), schema, compression="zstd")
    writer.write_table(table)
    return writer


def _find_sheet1_header_row(rows: list[tuple]) -> int | None:
    for i, row in enumerate(rows):
        if row and row[0] == "ป้ายชื่อแถว":
            return i
    return None


def _oracle_from_sheet1_ministry(rows: list[tuple], header_idx: int) -> dict:
    """2561-style: `Sheet1` แคบ 2 คอลัมน์ (ป้ายชื่อแถว/ผลรวม ของ พรบ.) รายกระทรวงทั้งหมด"""
    by_ministry: dict[str, float | None] = {}
    total: float | None = None
    for row in rows[header_idx + 1 :]:
        if not row or row[0] is None:
            continue
        name = clean(str(row[0]))
        value = _oracle_value(row[1]) if len(row) > 1 else None
        if name == "ผลรวมทั้งหมด":
            total = value
        else:
            by_ministry[name] = value
    return {"kind": "sheet1_ministry", "values": {"amount_thb": total}, "by_ministry": by_ministry}


# 2568 — Sheet1 pivot กว้าง กรองเฉพาะกระทรวงเดียว (ยืนยัน 02 §A1: สำนักนายกรัฐมนตรี)
PBO_2568_SHEET1_MINISTRY = "สำนักนายกรัฐมนตรี"

_SHEET1_TOTAL_COLUMNS: tuple[tuple[str, str, int], ...] = (
    ("current", "amount_thb", 1),
    ("current", "revised_thb", 2),
    ("current", "disbursed_thb", 3),
    ("capital", "amount_thb", 4),
    ("capital", "revised_thb", 5),
    ("capital", "disbursed_thb", 6),
    ("total", "amount_thb", 7),
    ("total", "revised_thb", 8),
    ("total", "disbursed_thb", 9),
)


def _oracle_from_sheet1_total(rows: list[tuple], header_idx: int) -> dict:
    """2568-style: `Sheet1` กว้าง header 2 ชั้น (ประจำ/ลงทุน × พรบ./หลังโอน/เบิกจ่าย)"""
    total_row = next(
        (row for row in rows[header_idx + 1 :] if row and row[0] == "ผลรวมทั้งหมด"), None
    )
    values: dict[str, dict[str, float | None]] = {"current": {}, "capital": {}, "total": {}}
    if total_row is not None:
        for split, field_name, col_idx in _SHEET1_TOTAL_COLUMNS:
            values[split][field_name] = (
                _oracle_value(total_row[col_idx]) if col_idx < len(total_row) else None
            )
    return {"kind": "sheet1_total", "values": values, "ministry_filter": PBO_2568_SHEET1_MINISTRY}


def _build_oracle(
    wb: openpyxl.Workbook,
    grand_total_values: dict[str, float | None] | None,
    n_rows_written: int,
    year: int,
) -> dict:
    if grand_total_values is not None:
        return {"kind": "grand_total", "values": grand_total_values}

    if "Sheet1" in wb.sheetnames:
        ws = wb["Sheet1"]
        rows = list(ws.iter_rows(values_only=True))
        header_idx = _find_sheet1_header_row(rows)
        if header_idx is not None:
            header_row = rows[header_idx]
            if len(header_row) <= 2:
                return _oracle_from_sheet1_ministry(rows, header_idx)
            return _oracle_from_sheet1_total(rows, header_idx)

    expected = YEAR_NO_ORACLE_EXPECTED_ROWS.get(year)
    return {"kind": "none", "n_rows": n_rows_written, "expected_rows": expected}


def _oracle_cache_path(cfg: PipelineConfig) -> Path:
    return cfg.cache_dir / "pbo" / ORACLE_FILENAME


def _headers_cache_path(cfg: PipelineConfig) -> Path:
    return cfg.cache_dir / "pbo" / HEADERS_FILENAME


def _merge_write_json(cfg: PipelineConfig, path: Path, year: int, entry: object) -> None:
    safe_path = cfg.assert_writable_path(path)
    safe_path.parent.mkdir(parents=True, exist_ok=True)
    data: dict = {}
    if safe_path.is_file():
        try:
            data = json.loads(safe_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            data = {}
    data[str(year)] = entry
    safe_path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
        newline="\n",
    )


def load_oracle(cfg: PipelineConfig) -> dict:
    path = _oracle_cache_path(cfg)
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def extract_year(
    cfg: PipelineConfig,
    year: int,
    path: Path,
    limit_rows: int | None = None,
) -> PboYearResult:
    """extract ไฟล์ PBO หนึ่งปี → `.cache/pbo/{year}.parquet` + อัปเดต oracle.json/headers.json"""
    start = time.monotonic()
    rel_path = path.relative_to(cfg.raw_data_dir).as_posix()
    source_doc_id = doc_id_for_path(rel_path)

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        sheet_name = _pick_data_sheet(wb.sheetnames, year, path)
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(max_col=MAX_COLS, values_only=True)
        header = next(rows_iter)
        _validate_header(header, year, path)

        is_partial = limit_rows is not None
        if not is_partial:
            # --limit-rows ต้องไม่แตะ headers.json/oracle.json ของการรันเต็ม (03 §4.1 แก้รอบ 2)
            _merge_write_json(
                cfg,
                _headers_cache_path(cfg),
                year,
                {
                    "sheet": sheet_name,
                    "header": [repr(h) for h in header],
                },
            )

        cache_filename = f"{year}{PARTIAL_SUFFIX}" if is_partial else f"{year}.parquet"
        cache_path = cfg.assert_writable_path(cfg.cache_dir / "pbo" / cache_filename)
        cache_path.parent.mkdir(parents=True, exist_ok=True)

        writer: pq.ParquetWriter | None = None
        buffer: list[dict] = []
        rows_written = 0
        n_grand_total_rows = 0
        grand_total_values: dict[str, float | None] | None = None
        flag_counts: Counter = Counter()

        excel_row = 1  # header คือแถว 1
        for row in rows_iter:
            excel_row += 1
            if _row_is_grand_total(row):
                n_grand_total_rows += 1
                if n_grand_total_rows > 1:
                    raise ValueError(
                        f"พบแถว 'Grand Total' มากกว่า 1 แถวใน {path} (ปี {year}, แถว Excel "
                        f"{excel_row}) — ไม่ตรงกับที่ยืนยันไว้ใน docs/02 ต้องตรวจไฟล์ใหม่"
                    )
                grand_total_values = {
                    field_name: _oracle_value(row[idx]) for idx, field_name, _label in MONEY_COLUMNS
                }
                continue

            if limit_rows is not None and rows_written >= limit_rows:
                # Grand Total (ถ้ามี) เจอก่อนเสมอ (แถวข้อมูลแรกสุด) จึงหยุด loop ได้ปลอดภัย
                # โดยไม่พลาดการตรวจ Grand Total — ประหยัดเวลาตอนทดสอบด้วย --limit-rows
                break

            record, flags = _build_record(row, year, excel_row, rel_path, sheet_name, source_doc_id)
            buffer.append(record)
            flag_counts.update(flags)
            rows_written += 1

            if len(buffer) >= CHUNK_ROWS:
                writer = _write_chunk(buffer, cache_path, writer)
                buffer.clear()

        if buffer:
            writer = _write_chunk(buffer, cache_path, writer)
        if writer is not None:
            writer.close()
        else:
            # ไม่มีแถวเลย (เช่น limit_rows=0) — เขียนไฟล์ parquet ว่างที่มี schema ถูกต้องไว้
            pq.write_table(pa.Table.from_pylist([], schema=pyarrow_schema()), str(cache_path))

        oracle = _build_oracle(wb, grand_total_values, rows_written, year)
        if not is_partial:
            _merge_write_json(cfg, _oracle_cache_path(cfg), year, oracle)
    finally:
        wb.close()

    cache_bytes = cache_path.stat().st_size
    elapsed = time.monotonic() - start
    return PboYearResult(
        year=year,
        file_path=path,
        rel_path=rel_path,
        sheet_name=sheet_name,
        rows_written=rows_written,
        n_grand_total_rows=n_grand_total_rows,
        flag_counts=flag_counts,
        cache_path=cache_path,
        cache_bytes=cache_bytes,
        elapsed_seconds=elapsed,
        oracle=oracle,
    )


def extract_all(
    cfg: PipelineConfig,
    years: list[int] | None = None,
    limit_rows: int | None = None,
) -> list[PboYearResult]:
    """extract หลายปี (ค่าเริ่มต้น = ทุกปีที่มีไฟล์จริง) เรียงตามปี"""
    year_files = discover_year_files(cfg, years)
    return [
        extract_year(cfg, year, path, limit_rows=limit_rows) for year, path in year_files.items()
    ]
