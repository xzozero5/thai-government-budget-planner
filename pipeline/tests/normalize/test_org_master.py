"""T-104: `normalize/org_master.py` — master กระทรวง/หน่วยงานจาก A2 + fuzzy matcher

ห้ามพึ่งไฟล์ raw จริง (สร้าง fixture xlsx ในเทสต์เอง) ยกเว้นเทสต์ที่ mark `@pytest.mark.rawdata`
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest
import yaml

from tgbp_pipeline.config import PipelineConfig
from tgbp_pipeline.normalize.org_master import (
    OrgMaster,
    OrgRecord,
    _code_str,
    _ensure_pseudo_ministries,
    build_org_master,
    export_catalog_orgs,
    load_org_master_cache,
)

A2_HEADER = ("min", "min_name", "agc", "agc_name")


def _write_a2_fixture(path: Path, rows: list[tuple]) -> None:
    """สร้าง xlsx ขนาดเล็กที่มีแค่ sheet `Data` (4 คอลัมน์ที่ org_master ใช้จริง)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(A2_HEADER)
    for row in rows:
        ws.append(row)
    wb.save(path)


def _write_config(tmp_path: Path) -> Path:
    pipeline_dir = tmp_path / "pipeline"
    pipeline_dir.mkdir()
    raw_dir = tmp_path / "raw"
    raw_dir.mkdir()
    (tmp_path / "web" / "public" / "data").mkdir(parents=True)
    (tmp_path / "web" / "tests" / "fixtures" / "data").mkdir(parents=True)

    config_path = pipeline_dir / "config.yaml"
    config_path.write_text(
        yaml.safe_dump(
            {
                "raw_data_dir": "../raw",
                "output_dir": "../web/public/data",
                "cache_dir": ".cache",
                "fixtures_dir": "../web/tests/fixtures/data",
            }
        ),
        encoding="utf-8",
    )
    return config_path


_BASIC_ROWS = [
    ("01000", "สำนักนายกรัฐมนตรี", "01001", "สำนักงานปลัดสำนักนายกรัฐมนตรี"),
    ("01000", "สำนักนายกรัฐมนตรี", "01001", "สำนักงานปลัดสำนักนายกรัฐมนตรี"),
    ("01000", "สำนักนายกรัฐมนตรี", "01002", "กรมประชาสัมพันธ์"),
    ("02000", "กระทรวงกลาโหม", "02001", "สำนักงานปลัดกระทรวงกลาโหม"),
]


# ---------------------------------------------------------------------------
# _code_str — เก็บรหัสเป็น string เสมอ (บาง agc ผสมตัวอักษร)
# ---------------------------------------------------------------------------


def test_code_str_keeps_leading_zero_string_as_is() -> None:
    assert _code_str("01000") == "01000"


def test_code_str_alnum_code_preserved() -> None:
    assert _code_str("7510A") == "7510A"


def test_code_str_none_is_empty_string() -> None:
    assert _code_str(None) == ""


def test_code_str_float_integer_from_excel_becomes_int_string() -> None:
    # openpyxl อาจอ่านเซลล์ตัวเลขล้วนเป็น float (เช่น 1000.0) — ต้องไม่มี ".0" หลุดมา
    assert _code_str(1000.0) == "1000"


# ---------------------------------------------------------------------------
# build_org_master — สร้าง master จาก fixture A2
# ---------------------------------------------------------------------------


def test_build_org_master_basic_records(tmp_path: Path) -> None:
    a2_path = tmp_path / "a2.xlsx"
    _write_a2_fixture(a2_path, _BASIC_ROWS)

    org_master = build_org_master(a2_path, write_cache=False)

    ministry_codes = {m.code for m in org_master.ministries if not m.synthetic}
    assert {"01000", "02000"} <= ministry_codes

    agency_codes = {a.code for a in org_master.agencies}
    assert agency_codes == {"01001", "01002", "02001"}

    agc_by_code = {a.code: a for a in org_master.agencies}
    assert agc_by_code["01001"].ministry_code == "01000"
    assert agc_by_code["01001"].name == "สำนักงานปลัดสำนักนายกรัฐมนตรี"
    assert agc_by_code["01001"].level == "agency"

    ministry_by_code = {m.code: m for m in org_master.ministries}
    assert ministry_by_code["01000"].level == "ministry"
    assert ministry_by_code["01000"].ministry_code is None


def test_build_org_master_picks_most_frequent_name_as_canonical(tmp_path: Path) -> None:
    a2_path = tmp_path / "a2.xlsx"
    rows = [
        ("80000", "กองทุนและเงินทุนหมุนเวียน", "80001", "กองทุน ก"),
        ("80000", "กองทุนและเงินทุนหมุนเวียน", "80001", "กองทุน ก"),
        ("80000", "ทุนหมุนเวียน", "80001", "กองทุน ก"),
    ]
    _write_a2_fixture(a2_path, rows)

    org_master = build_org_master(a2_path, write_cache=False)
    ministry = next(m for m in org_master.ministries if m.code == "80000")

    assert ministry.name == "กองทุนและเงินทุนหมุนเวียน"  # พบ 2 ครั้ง ชนะ "ทุนหมุนเวียน" (1 ครั้ง)
    assert "ทุนหมุนเวียน" in ministry.aliases  # ชื่อที่แพ้ถูกเก็บเป็น alias จากข้อมูลเอง


def test_build_org_master_agc_code_with_letters_preserved(tmp_path: Path) -> None:
    a2_path = tmp_path / "a2.xlsx"
    rows = [("75000", "องค์กรปกครองส่วนท้องถิ่น", "7510A", "เทศบาลตำบลเขาพนม")]
    _write_a2_fixture(a2_path, rows)

    org_master = build_org_master(a2_path, write_cache=False)
    assert org_master.agencies[0].code == "7510A"


def test_build_org_master_missing_required_column_raises(tmp_path: Path) -> None:
    a2_path = tmp_path / "a2_broken.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(("min", "min_name"))  # ขาด agc/agc_name
    ws.append(("01000", "สำนักนายกรัฐมนตรี"))
    wb.save(a2_path)

    with pytest.raises(ValueError, match="ขาดคอลัมน์"):
        build_org_master(a2_path, write_cache=False)


# ---------------------------------------------------------------------------
# org_aliases.yaml — โหลด alias มาต่อกับ master
# ---------------------------------------------------------------------------


def test_build_org_master_merges_yaml_ministry_alias(tmp_path: Path) -> None:
    a2_path = tmp_path / "a2.xlsx"
    _write_a2_fixture(a2_path, _BASIC_ROWS)

    alias_path = tmp_path / "aliases.yaml"
    alias_path.write_text(
        yaml.safe_dump(
            {
                "ministry": {"สำนักนายกรัฐมนตรี": ["สนร.", "สำนักนายก"]},
                "agency": {},
            },
            allow_unicode=True,
        ),
        encoding="utf-8",
    )

    org_master = build_org_master(a2_path, alias_path=alias_path, write_cache=False)
    ministry = next(m for m in org_master.ministries if m.code == "01000")
    assert "สนร." in ministry.aliases
    assert "สำนักนายก" in ministry.aliases

    result = org_master.match("สนร.", "สำนักงานปลัดสำนักนายกรัฐมนตรี")
    assert result.ministry_code == "01000"
    assert result.method == "alias"


def test_missing_alias_file_does_not_fail(tmp_path: Path) -> None:
    a2_path = tmp_path / "a2.xlsx"
    _write_a2_fixture(a2_path, _BASIC_ROWS)
    org_master = build_org_master(
        a2_path, alias_path=tmp_path / "does_not_exist.yaml", write_cache=False
    )
    assert len(org_master.ministries) >= 2


# ---------------------------------------------------------------------------
# _ensure_pseudo_ministries — synthetic code (X-xx) เฉพาะกลุ่มที่ A2 ไม่มีให้
# ---------------------------------------------------------------------------


def test_ensure_pseudo_ministries_no_synthetic_when_all_present() -> None:
    ministries = [
        OrgRecord(code=f"X{i:02d}", name=keyword, level="ministry", ministry_code=None)
        for i, keyword in enumerate(
            [
                "งบกลาง",
                "จังหวัดและกลุ่มจังหวัด",
                "รัฐวิสาหกิจ",
                "กองทุนและเงินทุนหมุนเวียน",
                "หน่วยงานของรัฐสภา",
                "หน่วยงานของศาล",
                "หน่วยงานขององค์กรอิสระและองค์กรอัยการ",
                "ส่วนราชการไม่สังกัดสำนักนายกรัฐมนตรี",
                "สภากาชาดไทย",
                "รายจ่ายเพื่อชดใช้เงินคงคลัง",
                "ส่วนราชการในพระองค์",
            ]
        )
    ]
    result = _ensure_pseudo_ministries(ministries)
    assert len(result) == len(ministries)  # ไม่มี synthetic เพิ่ม
    assert all(not m.synthetic for m in result)


def test_ensure_pseudo_ministries_adds_synthetic_when_missing() -> None:
    ministries = [
        OrgRecord(code="01000", name="สำนักนายกรัฐมนตรี", level="ministry", ministry_code=None)
    ]
    result = _ensure_pseudo_ministries(ministries)

    synthetic = [m for m in result if m.synthetic]
    assert len(synthetic) == 11  # ไม่มีกลุ่ม pseudo ไหนเลยใน input -> เพิ่มครบ 11
    assert all(m.code.startswith("X-") for m in synthetic)
    assert all(m.level == "ministry" for m in synthetic)
    # code synthetic ต้องไม่ปนกับ code จริง (prefix ต่างกันชัดเจน)
    real_codes = {m.code for m in ministries}
    assert not (real_codes & {m.code for m in synthetic})


# ---------------------------------------------------------------------------
# OrgMaster.match — exact / alias / fuzzy / none
# ---------------------------------------------------------------------------


@pytest.fixture
def org_master(tmp_path: Path) -> OrgMaster:
    a2_path = tmp_path / "a2.xlsx"
    rows = [
        *_BASIC_ROWS,
        ("07000", "กระทรวงเกษตรและสหกรณ์", "07001", "กรมชลประทาน"),
        ("07000", "กระทรวงเกษตรและสหกรณ์", "07002", "กรมประมง"),
    ]
    _write_a2_fixture(a2_path, rows)
    return build_org_master(a2_path, write_cache=False)


def test_match_exact(org_master: OrgMaster) -> None:
    result = org_master.match("สำนักนายกรัฐมนตรี", "กรมประชาสัมพันธ์")
    assert result.ministry_code == "01000"
    assert result.agency_code == "01002"
    assert result.method == "exact"
    assert result.score == 100.0
    assert result.flags == ()


def test_match_fuzzy_within_threshold(org_master: OrgMaster) -> None:
    # "กรมชลปะทาน" (พิมพ์ผิดตัดตัวอักษร) vs "กรมชลประทาน" จริง — ควรผ่าน threshold 92
    result = org_master.match("กระทรวงเกษตรและสหกรณ์", "กรมชลปะทาน")
    assert result.agency_code == "07001"
    assert result.method == "fuzzy"
    assert result.score >= 92.0


def test_match_below_threshold_returns_none_with_org_unmapped_flag(org_master: OrgMaster) -> None:
    result = org_master.match("กระทรวงเกษตรและสหกรณ์", "หน่วยงานที่ไม่มีอยู่จริงเลยสักนิด")
    assert result.agency_code is None
    assert result.method == "none"
    assert "org_unmapped" in result.flags


def test_match_unknown_ministry_and_agency_flags_both(org_master: OrgMaster) -> None:
    result = org_master.match("กระทรวงที่ไม่มีอยู่จริง", "หน่วยงานที่ไม่มีอยู่จริง")
    assert result.ministry_code is None
    assert result.agency_code is None
    assert "ministry_unmapped" in result.flags
    assert "org_unmapped" in result.flags


def test_match_fuzzy_restricted_to_same_ministry_no_cross_match(org_master: OrgMaster) -> None:
    """rapidfuzz candidate ต้องถูกจำกัดในกระทรวงที่ resolve แล้วเท่านั้น — ห้าม match ข้ามกระทรวง

    "กรมประมง" (07002, กระทรวงเกษตรฯ) สะกดใกล้กับ "กรมประชาสัมพันธ์" (01002, สนร.) ไม่พอที่จะ
    fuzzy ข้ามได้อยู่แล้ว แต่เทสต์นี้ยืนยันว่าแม้ผู้ใช้ระบุกระทรวงผิด (สนร.) ระบบไม่ไป "หยิบ"
    "กรมประมง" ของกระทรวงเกษตรฯ มาให้ (candidate ถูกจำกัดเฉพาะ agency ของกระทรวงที่ระบุ)
    """
    result = org_master.match("สำนักนายกรัฐมนตรี", "กรมประมง")
    assert result.agency_code is None
    assert "org_unmapped" in result.flags


def test_match_empty_inputs_return_none(org_master: OrgMaster) -> None:
    result = org_master.match(None, None)
    assert result.ministry_code is None
    assert result.agency_code is None
    assert result.method == "none"


def test_match_global_agency_fallback_when_ministry_unresolved(org_master: OrgMaster) -> None:
    # ไม่ระบุกระทรวง (หรือระบุผิดจนหาไม่เจอ) แต่ชื่อหน่วยงานตรงตัวและไม่กำกวมทั้ง master
    result = org_master.match(None, "กรมชลประทาน")
    assert result.agency_code == "07001"
    assert result.ministry_code == "07000"  # infer มาจาก agency ที่ resolve ได้
    assert "ministry_unmapped" not in result.flags


def test_match_local_gov_name_does_not_fuzzy_match_similar_name(tmp_path: Path) -> None:
    """T-105: ปิด fuzzy สำหรับชื่อ อปท. — "เทศบาลตำบลคลองโยง" ต้องไม่ match "...คลองยาง"

    (rapidfuzz WRatio ให้คะแนน 94.1 ซึ่งเกิน threshold 92 ปกติ แต่เป็นคนละที่กันจริง)
    """
    a2_path = tmp_path / "a2.xlsx"
    rows = [
        ("75000", "องค์กรปกครองส่วนท้องถิ่น", "7510A", "เทศบาลตำบลคลองยาง"),
    ]
    _write_a2_fixture(a2_path, rows)
    org_master = build_org_master(a2_path, write_cache=False)

    result = org_master.match("องค์กรปกครองส่วนท้องถิ่น", "เทศบาลตำบลคลองโยง")

    assert result.agency_code is None
    assert result.method == "none"
    assert "org_unmapped" in result.flags


def test_match_local_gov_name_still_matches_exact(tmp_path: Path) -> None:
    a2_path = tmp_path / "a2.xlsx"
    rows = [
        ("75000", "องค์กรปกครองส่วนท้องถิ่น", "7510A", "เทศบาลตำบลคลองโยง"),
    ]
    _write_a2_fixture(a2_path, rows)
    org_master = build_org_master(a2_path, write_cache=False)

    result = org_master.match("องค์กรปกครองส่วนท้องถิ่น", "เทศบาลตำบลคลองโยง")

    assert result.agency_code == "7510A"
    assert result.method == "exact"


def test_match_is_lru_cached(org_master: OrgMaster) -> None:
    org_master.match("สำนักนายกรัฐมนตรี", "กรมประชาสัมพันธ์")
    org_master.match("สำนักนายกรัฐมนตรี", "กรมประชาสัมพันธ์")
    info = org_master.cache_info()
    assert info.hits >= 1


# ---------------------------------------------------------------------------
# cache: write_org_master_cache / load_org_master_cache (N6 guard ผ่าน assert_writable_path)
# ---------------------------------------------------------------------------


def test_build_org_master_writes_cache_file(tmp_path: Path) -> None:
    config_path = _write_config(tmp_path)
    cfg = PipelineConfig.load(config_path)
    a2_path = tmp_path / "a2.xlsx"
    _write_a2_fixture(a2_path, _BASIC_ROWS)

    build_org_master(a2_path, cfg=cfg)

    cache_path = cfg.cache_dir / "org_master.json"
    assert cache_path.is_file()
    payload = json.loads(cache_path.read_text(encoding="utf-8"))
    assert "ministries" in payload
    assert "agencies" in payload
    assert any(m["code"] == "01000" for m in payload["ministries"])


def test_org_master_cache_uses_lf_line_endings(tmp_path: Path) -> None:
    config_path = _write_config(tmp_path)
    cfg = PipelineConfig.load(config_path)
    a2_path = tmp_path / "a2.xlsx"
    _write_a2_fixture(a2_path, _BASIC_ROWS)

    build_org_master(a2_path, cfg=cfg)

    raw_bytes = (cfg.cache_dir / "org_master.json").read_bytes()
    assert b"\r\n" not in raw_bytes


def test_cache_never_written_under_raw_dir(tmp_path: Path) -> None:
    """N6 guard: ถ้า cache_dir ถูกตั้งให้ชี้เข้า raw dir ต้อง raise ไม่เขียนทับข้อมูลดิบ"""
    pipeline_dir = tmp_path / "pipeline"
    pipeline_dir.mkdir()
    raw_dir = tmp_path / "raw"
    raw_dir.mkdir()
    (tmp_path / "web" / "public" / "data").mkdir(parents=True)
    (tmp_path / "web" / "tests" / "fixtures" / "data").mkdir(parents=True)

    config_path = pipeline_dir / "config.yaml"
    config_path.write_text(
        yaml.safe_dump(
            {
                "raw_data_dir": "../raw",
                "output_dir": "../web/public/data",
                "cache_dir": "../raw/should_not_write_here",
                "fixtures_dir": "../web/tests/fixtures/data",
            }
        ),
        encoding="utf-8",
    )
    cfg = PipelineConfig.load(config_path)

    from tgbp_pipeline.config import RawDataWriteError

    a2_path = tmp_path / "a2.xlsx"
    _write_a2_fixture(a2_path, _BASIC_ROWS)

    with pytest.raises(RawDataWriteError):
        build_org_master(a2_path, cfg=cfg)


def test_load_org_master_cache_roundtrip(tmp_path: Path) -> None:
    config_path = _write_config(tmp_path)
    cfg = PipelineConfig.load(config_path)
    a2_path = tmp_path / "a2.xlsx"
    _write_a2_fixture(a2_path, _BASIC_ROWS)

    original = build_org_master(a2_path, cfg=cfg)
    reloaded = load_org_master_cache(cfg)

    assert reloaded is not None
    assert {m.code for m in reloaded.ministries} == {m.code for m in original.ministries}
    assert {a.code for a in reloaded.agencies} == {a.code for a in original.agencies}

    # ผลลัพธ์ match ต้องเหมือนกันทุกประการหลังโหลดจาก cache
    assert reloaded.match("สำนักนายกรัฐมนตรี", "กรมประชาสัมพันธ์") == original.match(
        "สำนักนายกรัฐมนตรี", "กรมประชาสัมพันธ์"
    )


def test_load_org_master_cache_returns_none_when_missing(tmp_path: Path) -> None:
    config_path = _write_config(tmp_path)
    cfg = PipelineConfig.load(config_path)
    assert load_org_master_cache(cfg) is None


# ---------------------------------------------------------------------------
# export_catalog_orgs — สำหรับ T-110 (catalog/orgs.json)
# ---------------------------------------------------------------------------


def test_export_catalog_orgs_shape(tmp_path: Path) -> None:
    a2_path = tmp_path / "a2.xlsx"
    _write_a2_fixture(a2_path, _BASIC_ROWS)
    org_master = build_org_master(a2_path, write_cache=False)

    exported = export_catalog_orgs(org_master)
    assert len(exported) == len(org_master.ministries) + len(org_master.agencies)
    for entry in exported:
        assert set(entry.keys()) == {
            "code",
            "name",
            "level",
            "ministry_code",
            "aliases",
            "synthetic",
        }

    agency_entry = next(e for e in exported if e["code"] == "01001")
    assert agency_entry["ministry_code"] == "01000"
    assert agency_entry["level"] == "agency"


# ---------------------------------------------------------------------------
# rawdata: อัตรา map ของ distinct (กระทรวง, หน่วยงาน) PBO 2566 จริง เทียบ master จาก A2 จริง
# ---------------------------------------------------------------------------


def _real_a2_path(raw_data_dir: Path) -> Path:
    return (
        raw_data_dir
        / "งบประมาณ เชียงใหม่"
        / "1 - งบฟังก์ชันที่มาลงในเชียงใหม่"
        / "ร่าง พ.ร.บ. งบ 2570 ฉบับเต็ม - Excel.xlsx"
    )


def _iter_pbo_ministry_agency(path: Path):
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        sheet_name = next(s for s in wb.sheetnames if s.startswith("เบิกจ่ายภาพรวมทุกมิติ"))
        ws = wb[sheet_name]
        rows = ws.iter_rows(min_col=1, max_col=3, values_only=True)
        next(rows)  # header
        for _year, ministry, agency in rows:
            if ministry == "Grand Total":
                continue
            yield ministry, agency
    finally:
        wb.close()


@pytest.mark.rawdata
def test_rawdata_pbo_2566_org_mapping_rate() -> None:
    """T-104 DoD: ministry mapped >= 95%, agency mapped >= 80% ของ distinct pair จริงปี 2566"""
    cfg = PipelineConfig.load()
    a2_path = _real_a2_path(cfg.raw_data_dir)
    pbo_path = cfg.raw_data_dir / "PBO" / "2566.xlsx"
    if not a2_path.is_file() or not pbo_path.is_file():
        pytest.skip("ไม่มีโฟลเดอร์ข้อมูลดิบจริงบนเครื่องนี้")

    org_master = build_org_master(a2_path, write_cache=False)

    distinct_pairs = set(_iter_pbo_ministry_agency(pbo_path))
    assert len(distinct_pairs) > 0

    ministry_mapped = 0
    agency_mapped = 0
    for ministry, agency in distinct_pairs:
        result = org_master.match(ministry, agency)
        if result.ministry_code is not None:
            ministry_mapped += 1
        if result.agency_code is not None:
            agency_mapped += 1

    n = len(distinct_pairs)
    ministry_pct = ministry_mapped / n * 100
    agency_pct = agency_mapped / n * 100
    print(
        f"\nPBO 2566: {n} distinct (ministry, agency); "
        f"ministry={ministry_pct:.2f}% agency={agency_pct:.2f}%"
    )

    assert ministry_pct >= 95.0
    assert agency_pct >= 80.0
