"""T-107: `extract/local_sheets.py` — ข้อบัญญัติ/เทศบัญญัติ อปท. ปี 2570 (03 §4.3, 02 §A4)

ห้ามพึ่งไฟล์ raw จริง (สร้าง fixture xlsx เล็ก ๆ ในเทสต์เอง, เขียน output ลง `tmp_path` เท่านั้น
ห้ามแตะ `pipeline/.cache` จริง) ยกเว้นเทสต์ที่ mark `@pytest.mark.rawdata`
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pyarrow.parquet as pq
import pytest
import yaml

from tgbp_pipeline.config import PipelineConfig
from tgbp_pipeline.extract.local_sheets import (
    DATASET,
    RAJA_HEADER,
    V3GroupDiff,
    check_v3_raja,
    derive_province_and_local_gov,
    detect_generic_header,
    discover_local_files,
    extract_local_sheets,
    extract_one_file,
    extract_raja_summary_oracle,
    find_paired_pdfs,
    flag_amount_outliers,
    sheet_targets_fiscal_year_2570,
    slugify_ascii,
)
from tgbp_pipeline.util.hash import doc_id_for_path
from tgbp_pipeline.util.hash import source_id as compute_source_id

# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------


def _write_config(tmp_path: Path) -> Path:
    pipeline_dir = tmp_path / "pipeline"
    pipeline_dir.mkdir()
    (tmp_path / "raw").mkdir()
    (tmp_path / "web" / "public" / "data").mkdir(parents=True)
    (tmp_path / "web" / "tests" / "fixtures" / "data").mkdir(parents=True)

    config_path = pipeline_dir / "config.yaml"
    config_path.write_text(
        yaml.safe_dump(
            {
                "raw_data_dir": "../raw",
                "output_dir": "../web/public/data",
                "cache_dir": ".cache",
                "fixtures_dir": "../web/tests/fixtures/data",
            }
        ),
        encoding="utf-8",
    )
    return config_path


def _write_workbook(path: Path, sheets: dict[str, list[tuple]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)


def _read_cache_records(cache_path: Path) -> list[dict]:
    return pq.read_table(str(cache_path)).to_pylist()


def _by_item(records: list[dict], item_name: str) -> dict:
    return next(r for r in records if r["item_name_raw"] == item_name)


# ---------------------------------------------------------------------------
# discover_local_files / derive_province_and_local_gov / find_paired_pdfs
# ---------------------------------------------------------------------------


def test_discover_local_files_matches_pattern_and_excludes_act2570_files(tmp_path: Path) -> None:
    cfg = PipelineConfig.load(_write_config(tmp_path))
    folder = cfg.raw_data_dir / "งบประมาณ เชียงใหม่" / "2 - งบ อบจ. เชียงใหม่"
    folder.mkdir(parents=True)

    good = folder / "ร่างข้อบัญญัติงบ 2570 อบจ. เชียงใหม่ - Sheets.xlsx"
    good.write_bytes(b"x")
    # A3 (act2570_province) — ต้องไม่ถูกจับ แม้ลงท้าย " - Excel.xlsx" เหมือนกัน
    decoy = folder / "ร่าง พ.ร.บ. งบ 2570 เฉพาะเงินอุดหนุน อบจ. เชียงใหม่ - Excel.xlsx"
    decoy.write_bytes(b"x")
    # ไฟล์ไม่เกี่ยว
    (folder / "หมายเหตุ.xlsx").write_bytes(b"x")

    found = discover_local_files(cfg)
    assert found == [good]


def test_derive_province_and_local_gov_from_rel_path() -> None:
    rel_path = "งบประมาณ เชียงใหม่/2 - งบ อบจ. เชียงใหม่/ร่างข้อบัญญัติงบ 2570 อบจ. เชียงใหม่ - Sheets.xlsx"
    province, local_gov_name = derive_province_and_local_gov(rel_path)
    assert province == "เชียงใหม่"
    assert local_gov_name == "อบจ. เชียงใหม่"


def test_derive_province_and_local_gov_thesaban_naming() -> None:
    rel_path = "งบประมาณ เชียงใหม่/3 - งบเทศบาลนครเชียงใหม่/ร่างเทศบัญญัติงบ 2570 ทน. เชียงใหม่ - Excel.xlsx"
    province, local_gov_name = derive_province_and_local_gov(rel_path)
    assert province == "เชียงใหม่"
    assert local_gov_name == "ทน. เชียงใหม่"


def test_find_paired_pdfs_matches_title_stem_and_sorts(tmp_path: Path) -> None:
    cfg = PipelineConfig.load(_write_config(tmp_path))
    folder = cfg.raw_data_dir / "งบประมาณ เชียงใหม่" / "3 - งบเทศบาลนครเชียงใหม่"
    folder.mkdir(parents=True)

    xlsx_path = folder / "ร่างเทศบัญญัติงบ 2570 ทน. เชียงใหม่ - Excel.xlsx"
    xlsx_path.write_bytes(b"x")
    pdf1 = folder / "ร่างเทศบัญญัติงบ 2570 ทน. เชียงใหม่ - PDF เล่ม 1.pdf"
    pdf2 = folder / "ร่างเทศบัญญัติงบ 2570 ทน. เชียงใหม่ - PDF เล่ม 2.pdf"
    pdf1.write_bytes(b"x")
    pdf2.write_bytes(b"x")
    # ไฟล์ไม่เกี่ยว (คนละ title) ต้องไม่ถูกจับ
    (folder / "อื่น ๆ - PDF.pdf").write_bytes(b"x")

    found = find_paired_pdfs(xlsx_path, cfg)
    assert found == [pdf1, pdf2]


def test_slugify_ascii_is_ascii_and_deterministic() -> None:
    slug1 = slugify_ascii("อบต. ราชาเทวะ")
    slug2 = slugify_ascii("อบต. ราชาเทวะ")
    assert slug1 == slug2
    assert slug1.encode("ascii")  # ไม่ raise
    assert slugify_ascii("เชียงใหม่") != slugify_ascii("เชียงราย")


# ---------------------------------------------------------------------------
# sheet_targets_fiscal_year_2570
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("sheet_name", "expected"),
    [
        ("โครงการรวม งบ 70", True),
        ("โครงการรวม งบ 69", False),
        ("แผนพัฒฯ ปี 69", False),
        ("แผนพัฒฯ ปี 2569", False),
        ("ชีต1", True),  # ไม่มี token ปีเลย → เข้าเกณฑ์ default
        ("Data", True),
    ],
)
def test_sheet_targets_fiscal_year_2570(sheet_name: str, expected: bool) -> None:
    assert sheet_targets_fiscal_year_2570(sheet_name) is expected


# ---------------------------------------------------------------------------
# generic mapper: header detect สำเร็จ/ล้มเหลว + tie-break bug (REF_PDF_Page ว่าง vs
# REF_DOC_Page มีค่าจริง รูปแบบ "N/M") — regression ของบั๊กที่พบจริงใน ทน. เชียงใหม่ ชีต1
# ---------------------------------------------------------------------------


def test_detect_generic_header_success_samut_prakan_style() -> None:
    rows = [
        (None,) * 12,
        ("ร่างข้อบัญญัติงบประมาณรายจ่าย70",) + (None,) * 11,
        (
            "โครงการ",
            "แผนงาน ",
            "กลุ่มงาน",
            "ประเภทงบประมาณ",
            "ประเภทงบประมาณ (ย่อย)",
            "หน่วยรับงบประมาณ",
            "ราคา/หน่วย",
            "จำนวน (หน่วย)",
            "จำนวน (งวด)",
            "ยอดสุทธิ",
            "ร้อยละ",
            "แผนพัฒนาท้องถิ่น (2566-2570)",
        ),
        ("เงินเดือนนายก อบจ", "แผนงานบริหารทั่วไป", "งานบริหารทั่วไป", "งบบุคลากร") + (None,) * 8,
    ]
    result = detect_generic_header(rows)
    assert result is not None
    assert result.header_row_idx == 2
    assert result.mapping["item_name_raw"] == 0
    assert result.mapping["plan"] == 1
    assert result.mapping["amount_thb"] == 9


def test_detect_generic_header_fails_without_core_fields() -> None:
    rows = [
        ("ด้าน", "จำนวน", "ร้อยละ", "ส่วนต่าง"),
        ("แผนงานบริหารงานทั่วไป", 123, 0.5, -0.1),
    ]
    assert detect_generic_header(rows) is None


def test_generic_mapper_prefers_populated_page_column_over_empty_one_end_to_end(
    tmp_path: Path,
) -> None:
    """Regression: `REF_PDF_Page` (ว่างทั้งไฟล์) กับ `REF_DOC_Page` (มีค่าจริง รูปแบบ `"N/M"`)
    แข่งกันเป็น `source_page` คะแนนเท่ากัน — ต้องเลือก `REF_DOC_Page` (ไม่ใช่คอลัมน์แรกที่ว่าง)
    และแปลง `"5/99"` → หน้า `5` ได้ (03 §4.3, ยืนยันจริงจาก `ทน. เชียงใหม่` sheet `ชีต1`)
    """
    cfg = PipelineConfig.load(_write_config(tmp_path))
    header = (
        "REF_PDF_Page",
        "REF_DOC_Page",
        "Divison",
        "Budgetary_Plan",
        "Budgetary_USER",
        "Budgetary_TYPE",
        "Budgetary_TYPE2",
        "Pre Project",
        "Project ",
        "number",
        "amount",
        "FCY",
    )
    data_row = (
        None,
        "5/99",
        "กองการเจ้าหน้าที่",
        "แผนงานงบกลาง",
        "งบกลาง",
        "งบกลาง",
        "งบกลาง",
        "เงินสมทบกองทุนประกันสังคม",
        "เงินสมทบกองทุนประกันสังคม",
        None,
        9514200,
        None,
    )
    folder = cfg.raw_data_dir / "งบประมาณ เชียงใหม่" / "3 - งบเทศบาลนครเชียงใหม่"
    xlsx_path = folder / "ร่างเทศบัญญัติงบ 2570 ทน. เชียงใหม่ - Excel.xlsx"
    _write_workbook(xlsx_path, {"ชีต1": [header, data_row]})

    local_dir = tmp_path / "cache" / "local"
    local_dir.mkdir(parents=True)
    result = extract_one_file(cfg, xlsx_path, local_dir)

    assert result.sheets_used == ["ชีต1"]
    records = _read_cache_records(result.cache_path)
    assert len(records) == 1
    assert records[0]["source_page"] == 5  # จาก "5/99" ไม่ใช่ None จาก REF_PDF_Page
    assert records[0]["amount_thb"] == 9_514_200
    assert records[0]["agency"] == "กองการเจ้าหน้าที่"  # มาจาก Divison (typo) ผ่าน generic mapper


# ---------------------------------------------------------------------------
# ราชาเทวะ — mapper เฉพาะ + V3
# ---------------------------------------------------------------------------


def _raja_row(
    page: int,
    plan: str,
    work: str | None,
    budget_group: str,
    item: str,
    amount: object,
    *,
    expense_category: str | None = None,
    sub_category: str | None = None,
    department: str | None = "กองคลัง",
    description: str | None = "รายละเอียด",
    legal_reference: str | None = "ระเบียบ ก",
) -> tuple:
    return (
        page,
        plan,
        work,
        budget_group,
        expense_category,
        sub_category,
        item,
        amount,
        department,
        description,
        legal_reference,
    )


def _build_raja_workbook(tmp_path: Path) -> tuple[PipelineConfig, Path]:
    cfg = PipelineConfig.load(_write_config(tmp_path))
    folder = cfg.raw_data_dir / "งบประมาณ สมุทรปราการ" / "3 - งบ อบต. ราชาเทวะ"
    xlsx_path = folder / "ร่างข้อบัญญัติงบ 2570 อบต. ราชาเทวะ - Sheets.xlsx"

    plan_rows = [
        RAJA_HEADER,
        _raja_row(10, "แผนงานทดสอบ", "งานเอ", "งบดำเนินงาน", "ค่าซ่อมแซม", 50_000),
        _raja_row(11, "แผนงานทดสอบ", "งานบี", "งบบุคลากร", "เงินเดือน", 200_000),
        # outlier จริง — จำลองเคส "เงินสมทบกองทุนประกันสังคม" amount=2 ใน 02 §A4
        _raja_row(
            12,
            "แผนงานทดสอบ",
            "งานบี",
            "งบบุคลากร",
            "เงินสมทบกองทุนประกันสังคม",
            2,
        ),
    ]
    summary_rows = [
        ("page", "plan", "work", "budget_group", "total_amount"),
        (10, "แผนงานทดสอบ", "งานเอ", "งบดำเนินงาน", 50_000),
        # oracle ของกลุ่มนี้ตรงกับผลรวมแค่แถวที่ไม่ใช่ outlier (200,000) — ยืนยันว่า check_v3_raja
        # ต้องตัดแถว outlier (amount=2) ออกจากผลรวมก่อนเทียบ
        (11, "แผนงานทดสอบ", "งานบี", "งบบุคลากร", 200_000),
    ]
    _write_workbook(
        xlsx_path,
        {
            "แผนงานทดสอบ": plan_rows,
            "ocr_raw_data": [RAJA_HEADER],
            "pivot_ocr_raw_data": [("x",)],
            "summary_ocr_raw_data": summary_rows,
        },
    )
    return cfg, xlsx_path


def test_raja_extraction_flags_upstream_ocr_and_outlier_v3_passes(tmp_path: Path) -> None:
    cfg, xlsx_path = _build_raja_workbook(tmp_path)
    local_dir = tmp_path / "cache" / "local"
    local_dir.mkdir(parents=True)

    result = extract_one_file(cfg, xlsx_path, local_dir)

    assert result.sheets_used == ["แผนงานทดสอบ"]
    assert {s.sheet_name for s in result.sheets_skipped} == {
        "ocr_raw_data",
        "pivot_ocr_raw_data",
        "summary_ocr_raw_data",
    }
    records = _read_cache_records(result.cache_path)
    assert len(records) == 3
    assert all("upstream_ocr" in r["quality_flags"] for r in records)

    outlier = _by_item(records, "เงินสมทบกองทุนประกันสังคม")
    assert outlier["amount_thb"] == 2
    assert "amount_outlier" in outlier["quality_flags"]

    normal = _by_item(records, "เงินเดือน")
    assert "amount_outlier" not in normal["quality_flags"]

    assert result.v3 is not None
    assert result.v3.n_outlier_rows_excluded == 1
    assert result.v3.passed is True


def test_raja_v3_fails_and_reports_diff_when_mismatched(tmp_path: Path) -> None:
    cfg, xlsx_path = _build_raja_workbook(tmp_path)
    # แก้ summary ให้ไม่ตรงกับ plan sheet (จงใจ mismatch กลุ่ม "งานเอ")
    wb = openpyxl.load_workbook(str(xlsx_path))
    ws = wb["summary_ocr_raw_data"]
    ws["E2"] = 999_999  # เดิม 50,000
    wb.save(xlsx_path)

    local_dir = tmp_path / "cache" / "local"
    local_dir.mkdir(parents=True)
    result = extract_one_file(cfg, xlsx_path, local_dir)

    assert result.v3 is not None
    assert result.v3.passed is False
    diff_keys = {d.key for d in result.v3.diffs}
    assert ("แผนงานทดสอบ", "งานเอ", "งบดำเนินงาน") in diff_keys
    bad = next(d for d in result.v3.diffs if d.key == ("แผนงานทดสอบ", "งานเอ", "งบดำเนินงาน"))
    assert bad.computed_baht == 50_000
    assert bad.oracle_baht == 999_999.0


def test_check_v3_raja_direct_pass_and_fail() -> None:
    records = [
        {
            "plan": "P",
            "activity": "A",
            "budget_type": "B",
            "amount_thb": 100,
            "quality_flags": [],
        },
        {
            "plan": "P",
            "activity": "A",
            "budget_type": "B",
            "amount_thb": 5,
            "quality_flags": ["amount_outlier"],
        },
    ]
    oracle_ok = {("P", "A", "B"): 100.0}
    result_ok = check_v3_raja(records, oracle_ok)
    assert result_ok.passed is True
    assert result_ok.n_outlier_rows_excluded == 1

    oracle_bad = {("P", "A", "B"): 999.0}
    result_bad = check_v3_raja(records, oracle_bad)
    assert result_bad.passed is False
    assert result_bad.diffs == [
        V3GroupDiff(key=("P", "A", "B"), computed_baht=100, oracle_baht=999.0, diff=-899.0)
    ]


def test_extract_raja_summary_oracle_sums_duplicate_keys() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(("page", "plan", "work", "budget_group", "total_amount"))
    ws.append((1, "P", "A", "B", 10))
    ws.append((2, "P", "A", "B", 5))  # คีย์ซ้ำ ต้องบวกกัน
    oracle = extract_raja_summary_oracle(ws)
    assert oracle == {("P", "A", "B"): 15.0}


# ---------------------------------------------------------------------------
# amount outlier flag — เฉพาะ absolute threshold (regression: z-score เดิมทำให้ยอดถูกต้อง
# หลักสิบล้านบาทถูกแฟล็กผิดเพราะกลุ่ม expense_category=None ปนกันหมด)
# ---------------------------------------------------------------------------


def test_flag_amount_outliers_absolute_threshold_only_no_false_positive_on_large_legit_amount() -> (
    None
):
    records = [
        {"amount_thb": 2, "expense_category": None, "quality_flags": []},
        {"amount_thb": 80_000, "expense_category": None, "quality_flags": []},
        {"amount_thb": 48_000_000, "expense_category": None, "quality_flags": []},
        {"amount_thb": 14_000_000, "expense_category": None, "quality_flags": []},
        {"amount_thb": 100_000, "expense_category": None, "quality_flags": []},
        {"amount_thb": None, "expense_category": None, "quality_flags": []},
    ]
    flag_amount_outliers(records)
    flagged = [r["amount_thb"] for r in records if "amount_outlier" in r["quality_flags"]]
    assert flagged == [2]


def test_flag_amount_outliers_does_not_duplicate_flag() -> None:
    records = [{"amount_thb": 1, "expense_category": None, "quality_flags": ["amount_outlier"]}]
    flag_amount_outliers(records)
    assert records[0]["quality_flags"] == ["amount_outlier"]


# ---------------------------------------------------------------------------
# อบจ. เชียงใหม่ "Data" sheet — Divison (typo, ว่าง) vs Division (มีค่า) + FCY → description
# ---------------------------------------------------------------------------


def test_cnx_pcao_data_sheet_prefers_division_over_divison_typo(tmp_path: Path) -> None:
    cfg = PipelineConfig.load(_write_config(tmp_path))
    folder = cfg.raw_data_dir / "งบประมาณ เชียงใหม่" / "2 - งบ อบจ. เชียงใหม่"
    xlsx_path = folder / "ร่างข้อบัญญัติงบ 2570 อบจ. เชียงใหม่ - Sheets.xlsx"

    header = (
        "REF_DOC_Page",
        "Divison",
        "Budgetary_Plan",
        "Budgetary_USER",
        "Budgetary_TYPE",
        "Budgetary_TYPE2",
        "Pre Project",
        "Project ",
        "number",
        "amount",
        "Division",
        "FCY",
    )
    row_with_division = (
        260.0,
        None,
        "แผนงานบริหารงานทั่วไป",
        "งานบริหารทั่วไป",
        "งบบุคลากร",
        "เงินเดือน",
        "-",
        "เงินเดือนพนักงาน",
        1,
        500_000,
        "กองคลัง",
        "หมายเหตุ FCY",
    )
    # เคสสมมติ: Division ว่าง แต่ Divison (typo) มีค่า → fallback ต้องทำงาน + ติด flag
    row_fallback_divison = (
        261.0,
        "กองช่างทดแทน",
        "แผนงานบริหารงานทั่วไป",
        "งานบริหารทั่วไป",
        "งบดำเนินงาน",
        "ค่าใช้สอย",
        "-",
        "ค่าจ้างเหมาบริการ",
        1,
        20_000,
        None,
        "รถ 9 คัน",
    )
    _write_workbook(xlsx_path, {"Data": [header, row_with_division, row_fallback_divison]})

    local_dir = tmp_path / "cache" / "local"
    local_dir.mkdir(parents=True)
    result = extract_one_file(cfg, xlsx_path, local_dir)

    assert result.sheets_used == ["Data"]
    records = _read_cache_records(result.cache_path)
    assert len(records) == 2

    r1 = _by_item(records, "เงินเดือนพนักงาน")
    assert r1["agency"] == "กองคลัง"
    assert r1["source_page"] == 260
    assert r1["description"] == "หมายเหตุ FCY"  # FCY ไม่ใช่ปีงบ → description
    assert "agency_from_divison_typo_column" not in r1["quality_flags"]

    r2 = _by_item(records, "ค่าจ้างเหมาบริการ")
    assert r2["agency"] == "กองช่างทดแทน"
    assert "agency_from_divison_typo_column" in r2["quality_flags"]
    assert r2["description"] == "รถ 9 คัน"


# ---------------------------------------------------------------------------
# อบจ. สมุทรปราการ — หลาย sheet คนละทรง ต้องเลือกใช้แค่ sheet เดียว (โครงการรวม งบ 70)
# ---------------------------------------------------------------------------


def test_samut_prakan_multi_sheet_selects_only_fiscal_2570_sheet(tmp_path: Path) -> None:
    cfg = PipelineConfig.load(_write_config(tmp_path))
    folder = cfg.raw_data_dir / "งบประมาณ สมุทรปราการ" / "2 - งบ อบจ. สมุทรปราการ"
    xlsx_path = folder / "ร่างข้อบัญญัติงบ 2570 อบจ. สมุทรปราการ - Sheets.xlsx"

    overview_rows = [
        (None, "งบปี 2567", None, "งบปี 2570"),
        ("ด้าน ", "จำนวน", "ร้อยละ", "จำนวน"),
        ("แผนงานบริหารงานทั่วไป", 100.0, 0.5, 200.0),
    ]
    project_header = (
        "โครงการ",
        "แผนงาน ",
        "กลุ่มงาน",
        "ประเภทงบประมาณ",
        "ประเภทงบประมาณ (ย่อย)",
        "หน่วยรับงบประมาณ",
        "ราคา/หน่วย",
        "จำนวน (หน่วย)",
        "จำนวน (งวด)",
        "ยอดสุทธิ",
        "ร้อยละ",
        "แผนพัฒนาท้องถิ่น (2566-2570)",
    )
    rows_70 = [
        (None,) * 12,
        ("ร่างข้อบัญญัติงบประมาณรายจ่าย70",) + (None,) * 11,
        project_header,
        (
            "เงินเดือนนายก อบจ",
            "แผนงานบริหารทั่วไป",
            "งานบริหารทั่วไป",
            "งบบุคลากร (ฝ่ายการเมือง)",
            "งบบุคลากร (ฝ่ายการเมือง)",
            "สำนักงานเลขา",
            55_530.0,
            1.0,
            12.0,
            666_360,
            None,
            None,
        ),
    ]
    rows_69 = [
        (None,) * 11,
        ("โครงการรวมงบปี 2569",) + (None,) * 10,
        project_header[:11],
        (
            "เงินเดือนนายก อบจ",
            "แผนงานบริหารทั่วไป",
            "งานบริหารทั่วไป",
            "งบบุคลากร (ฝ่ายการเมือง)",
            "งบบุคลากร (ฝ่ายการเมือง)",
            None,
            55_530.0,
            1.0,
            12.0,
            666_360,
            None,
        ),
    ]
    _write_workbook(
        xlsx_path,
        {
            "ภาพรวม": overview_rows,
            "โครงการรวม งบ 70": rows_70,
            "โครงการรวม งบ 69": rows_69,
        },
    )

    local_dir = tmp_path / "cache" / "local"
    local_dir.mkdir(parents=True)
    result = extract_one_file(cfg, xlsx_path, local_dir)

    assert result.sheets_used == ["โครงการรวม งบ 70"]
    skip_reasons = {s.sheet_name: s.reason for s in result.sheets_skipped}
    assert "ภาพรวม" in skip_reasons
    assert "map ไม่ได้" in skip_reasons["ภาพรวม"]
    assert "ระบุปีงบอื่น" in skip_reasons["โครงการรวม งบ 69"]

    records = _read_cache_records(result.cache_path)
    assert len(records) == 1
    assert records[0]["amount_thb"] == 666_360
    assert records[0]["plan"] == "แผนงานบริหารทั่วไป"


# ---------------------------------------------------------------------------
# source_id golden + stability
# ---------------------------------------------------------------------------


def test_raja_source_id_stable_and_matches_hash_formula(tmp_path: Path) -> None:
    cfg, xlsx_path = _build_raja_workbook(tmp_path)
    rel_path = xlsx_path.relative_to(cfg.raw_data_dir).as_posix()

    local_dir1 = tmp_path / "cache1" / "local"
    local_dir1.mkdir(parents=True)
    result1 = extract_one_file(cfg, xlsx_path, local_dir1)
    records1 = _read_cache_records(result1.cache_path)

    local_dir2 = tmp_path / "cache2" / "local"
    local_dir2.mkdir(parents=True)
    result2 = extract_one_file(cfg, xlsx_path, local_dir2)
    records2 = _read_cache_records(result2.cache_path)

    ids1 = sorted(r["source_id"] for r in records1)
    ids2 = sorted(r["source_id"] for r in records2)
    assert ids1 == ids2

    first_row = _by_item(records1, "ค่าซ่อมแซม")
    expected_id = compute_source_id(DATASET, rel_path, "แผนงานทดสอบ", 2)
    assert first_row["source_id"] == expected_id

    expected_doc_id = doc_id_for_path(rel_path)
    assert all(r["source_doc_id"] == expected_doc_id for r in records1)


# ---------------------------------------------------------------------------
# extract_local_sheets — orchestration ทั้งไฟล์ + pdf_pairs.json
# ---------------------------------------------------------------------------


def test_extract_local_sheets_end_to_end_writes_pdf_pairs(tmp_path: Path) -> None:
    cfg, xlsx_path = _build_raja_workbook(tmp_path)
    pdf_path = xlsx_path.parent / "ร่างข้อบัญญัติงบ 2570 อบต. ราชาเทวะ - PDF.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 fake")

    cache_dir = tmp_path / "cache"
    report = extract_local_sheets(cfg, cache_dir=cache_dir, files=[xlsx_path])

    assert report.total_files == 1
    assert report.total_rows == 3
    assert report.pdf_pairs_path is not None
    assert report.pdf_pairs_path.is_relative_to(cache_dir)

    only_result = report.results[0]
    assert only_result.province == "สมุทรปราการ"
    assert only_result.local_gov_name == "อบต. ราชาเทวะ"
    assert only_result.cache_path.is_relative_to(cache_dir)
    assert only_result.cache_bytes > 0

    import json

    pairs = json.loads(report.pdf_pairs_path.read_text(encoding="utf-8"))
    assert only_result.source_doc_id in pairs
    entry = pairs[only_result.source_doc_id]
    assert entry["pdf_rel_paths"] == [
        "งบประมาณ สมุทรปราการ/3 - งบ อบต. ราชาเทวะ/ร่างข้อบัญญัติงบ 2570 อบต. ราชาเทวะ - PDF.pdf"
    ]


def test_extract_local_sheets_no_pdf_pairs_file_when_no_matches(tmp_path: Path) -> None:
    cfg, xlsx_path = _build_raja_workbook(tmp_path)
    # ไม่สร้างไฟล์ PDF คู่กัน
    cache_dir = tmp_path / "cache"
    report = extract_local_sheets(cfg, cache_dir=cache_dir, files=[xlsx_path])
    assert report.pdf_pairs_path is None


# ---------------------------------------------------------------------------
# rawdata: รัน extract_local_sheets จริงกับไฟล์จริง (ห้ามแตะ pipeline/.cache จริง)
# ---------------------------------------------------------------------------


@pytest.mark.rawdata
def test_rawdata_extract_local_sheets_runs_on_real_files(tmp_path: Path) -> None:
    real_cfg = PipelineConfig.load()
    files = discover_local_files(real_cfg)
    if not files:
        pytest.skip("ไม่มีโฟลเดอร์ข้อมูลดิบจริงบนเครื่องนี้")
    assert len(files) == 4

    cache_dir = tmp_path / "cache"
    report = extract_local_sheets(real_cfg, cache_dir=cache_dir)
    assert cache_dir.is_relative_to(tmp_path)
    assert report.total_files == 4
    assert report.total_rows > 0
    for result in report.results:
        assert result.cache_path.is_relative_to(tmp_path)
