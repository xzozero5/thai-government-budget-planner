"""T-108: `extract/committee_xlsx.py` — เอกสาร กมธ.ติดตามงบ Excel/XLS (A5, 03 §4.4)

ห้ามพึ่งไฟล์ raw จริง (สร้าง fixture xlsx เล็ก ๆ ในเทสต์เอง, เขียน output ลง `tmp_path` เท่านั้น
ห้ามแตะ `pipeline/.cache` จริง) ยกเว้นเทสต์ที่ mark `@pytest.mark.rawdata`
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pyarrow.parquet as pq
import pytest
import yaml

from tgbp_pipeline.config import PipelineConfig
from tgbp_pipeline.extract.committee_xlsx import (
    COMMITTEE_ROOT,
    DATASET,
    UnsupportedWorkbookError,
    _agency_from_rel_path,
    _extract_fiscal_year_be,
    _gov_level_for_agency,
    _is_subtotal_row,
    _is_zip_magic,
    _iter_workbook_sheets,
    detect_table_region,
    detect_unit,
    discover_committee_files,
    extract_committee_file,
    map_columns,
    process_sheet,
)
from tgbp_pipeline.extract.office_text import read_doc_chunks_gz
from tgbp_pipeline.util.hash import doc_id_for_path
from tgbp_pipeline.util.hash import source_id as compute_source_id

# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------


def _write_config(tmp_path: Path) -> Path:
    pipeline_dir = tmp_path / "pipeline"
    pipeline_dir.mkdir()
    (tmp_path / "raw").mkdir()
    (tmp_path / "web" / "public" / "data").mkdir(parents=True)
    (tmp_path / "web" / "tests" / "fixtures" / "data").mkdir(parents=True)

    config_path = pipeline_dir / "config.yaml"
    config_path.write_text(
        yaml.safe_dump(
            {
                "raw_data_dir": "../raw",
                "output_dir": "../web/public/data",
                "cache_dir": ".cache",
                "fixtures_dir": "../web/tests/fixtures/data",
            }
        ),
        encoding="utf-8",
    )
    return config_path


def _write_workbook(path: Path, sheets: dict[str, list[list]]) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)


# แผนงาน "ปี 63" ทั่วไปที่ map ได้: title (2 แถว) + blank + header + sub-header หน่วย + ข้อมูล
def _mapped_sheet_rows(unit_hint: str = "(บาท)") -> list[list]:
    return [
        ["กองทุนทดสอบ", None, None, None, None],
        ["รายละเอียดโครงการปีงบประมาณ พ.ศ. 2563", None, None, None, None],
        [None, None, None, None, None],
        ["ลำดับ", "ชื่อโครงการ", "หน่วยงาน", "จังหวัด", "งบประมาณ"],
        [None, None, None, None, unit_hint],
        [1, "โครงการก่อสร้างฝายน้ำล้น", "เทศบาลตำบลบ้านโป่ง", "เชียงใหม่", 66000],
        [2, "โครงการขุดบ่อบาดาล", "กรมชลประทาน", "ขอนแก่น", 120000],
        [None, "รวม", None, None, 186000],
    ]


# ---------------------------------------------------------------------------
# detect_table_region / map_columns / detect_unit
# ---------------------------------------------------------------------------


def test_detect_table_region_header_not_in_first_row() -> None:
    rows = _mapped_sheet_rows()
    region = detect_table_region(rows)
    assert region is not None
    assert region.header_row_idx == 3  # ไม่ใช่แถวแรก — มี title 2 แถว + แถวว่างนำหน้า
    assert region.data_start_idx == 5
    assert region.data_end_idx == 8


def test_detect_table_region_none_when_no_table() -> None:
    rows = [["หมายเหตุ"], [None], ["ไม่มีตารางในชีตนี้"]]
    assert detect_table_region(rows) is None


def test_map_columns_maps_required_and_optional_fields() -> None:
    header = ("ลำดับ", "ชื่อโครงการ", "หน่วยงาน", "จังหวัด", "งบประมาณ")
    mapping = map_columns(header)
    assert mapping["item_name"] == 1
    assert mapping["amount_thb"] == 4
    assert mapping["agency"] == 2
    assert mapping["province"] == 3


def test_detect_unit_baht() -> None:
    rows = _mapped_sheet_rows(unit_hint="(บาท)")
    region = detect_table_region(rows)
    assert detect_unit(rows, region) == ("บาท", 1)


def test_detect_unit_million_baht() -> None:
    rows = _mapped_sheet_rows(unit_hint="(ล้านบาท)")
    region = detect_table_region(rows)
    assert detect_unit(rows, region) == ("ล้านบาท", 1_000_000)


def test_detect_unit_thousand_baht() -> None:
    rows = _mapped_sheet_rows(unit_hint="(พันบาท)")
    region = detect_table_region(rows)
    assert detect_unit(rows, region) == ("พันบาท", 1_000)


def test_detect_unit_unknown_returns_none() -> None:
    rows = [
        ["กองทุนทดสอบ", None, None, None],
        ["ลำดับ", "รายการ", "หน่วยงาน", "จำนวนเงิน"],
        [1, "รายการทดสอบ", "กรมทดสอบ", 5000],
    ]
    region = detect_table_region(rows)
    assert region is not None
    assert detect_unit(rows, region) is None


# ---------------------------------------------------------------------------
# process_sheet: map ได้ / map ไม่ได้ (ไม่รู้หน่วยเงิน) / แถว subtotal
# ---------------------------------------------------------------------------


def test_process_sheet_maps_amount_and_year_and_gov_level_per_row() -> None:
    rows = _mapped_sheet_rows()
    result = process_sheet(rows, "ปี 63", rel_path="กมธ.ติดตามงบ/x/y/z.xlsx", source_doc_id="d_x")

    assert result.mapped is True
    assert result.n_rows == 2  # แถว "รวม" ถูกตัดออก ไม่นับเป็นรายการ
    assert result.subtotal_rows_skipped == 1

    by_agency = {r["agency"]: r for r in result.records}
    fay = by_agency["เทศบาลตำบลบ้านโป่ง"]
    kaset = by_agency["กรมชลประทาน"]

    # หน่วย "บาท" (multiplier=1) → amount ตรงตัว
    assert fay["amount_thb"] == 66000
    assert kaset["amount_thb"] == 120000

    # ปีงบต้องมาจากชื่อ sheet/title ("ปีงบประมาณ พ.ศ. 2563") ทุกแถวในชีตเดียวกัน
    assert fay["fiscal_year_be"] == 2563
    assert kaset["fiscal_year_be"] == 2563
    assert fay["fiscal_year_ce"] == 2020
    assert fay["quality_flags"] == []

    # gov_level ต้องต่อแถวจริง (เทศบาล = local, กรม = central) — ไม่ใช่ค่าคงที่ทั้ง sheet
    assert fay["gov_level"] == "local"
    assert kaset["gov_level"] == "central"


def test_process_sheet_million_baht_multiplies_correctly() -> None:
    rows = _mapped_sheet_rows(unit_hint="(ล้านบาท)")
    result = process_sheet(rows, "ปี 63", rel_path="กมธ.ติดตามงบ/x/y/z.xlsx", source_doc_id="d_x")
    by_agency = {r["agency"]: r for r in result.records}
    assert by_agency["เทศบาลตำบลบ้านโป่ง"]["amount_thb"] == 66_000_000_000
    assert by_agency["เทศบาลตำบลบ้านโป่ง"]["amount_unit_source"] == "ล้านบาท"


def test_process_sheet_thousand_baht_multiplies_correctly() -> None:
    rows = _mapped_sheet_rows(unit_hint="(พันบาท)")
    result = process_sheet(rows, "ปี 63", rel_path="กมธ.ติดตามงบ/x/y/z.xlsx", source_doc_id="d_x")
    by_agency = {r["agency"]: r for r in result.records}
    assert by_agency["เทศบาลตำบลบ้านโป่ง"]["amount_thb"] == 66_000_000
    assert by_agency["เทศบาลตำบลบ้านโป่ง"]["amount_unit_source"] == "พันบาท"


def test_process_sheet_unknown_unit_not_mapped() -> None:
    """N3: ไม่รู้หน่วยเงิน → ห้ามเดา → ไม่เข้า budget_lines (mapped=False)"""
    rows = [
        ["กองทุนทดสอบ", None, None, None],
        ["ลำดับ", "รายการ", "หน่วยงาน", "จำนวนเงิน"],
        [1, "รายการทดสอบ", "กรมทดสอบ", 5000],
    ]
    result = process_sheet(rows, "mask1", rel_path="กมธ.ติดตามงบ/x/y/z.xlsx", source_doc_id="d_x")
    assert result.mapped is False
    assert result.records == []
    assert "ไม่รู้หน่วยเงิน" in (result.reason or "")


def test_process_sheet_missing_required_field_not_mapped() -> None:
    rows = [
        ["ลำดับ", "หน่วยงาน", "จังหวัด", "หมายเหตุ"],
        [1, "กรมทดสอบ", "เชียงใหม่", "ไม่มี item_name/amount ให้ map"],
    ]
    result = process_sheet(rows, "sheet1", rel_path="กมธ.ติดตามงบ/x/y/z.xlsx", source_doc_id="d_x")
    assert result.mapped is False
    assert "map คอลัมน์ไม่ครบ" in (result.reason or "")


def test_process_sheet_only_subtotal_rows_not_mapped() -> None:
    """ถ้าตัดแถว subtotal แล้วไม่เหลือรายการจริงเลย → mapped=False (ไม่ใช่ 0-row budget_lines)"""
    rows = [
        ["ลำดับ", "ชื่อโครงการ", "หน่วยงาน", "งบประมาณ (บาท)"],
        [None, "รวมทั้งสิ้น", None, 100000],
    ]
    result = process_sheet(rows, "sheet1", rel_path="กมธ.ติดตามงบ/x/y/z.xlsx", source_doc_id="d_x")
    assert result.mapped is False
    assert result.subtotal_rows_skipped == 1
    assert "ยอดรวม" in (result.reason or "")


def test_is_subtotal_row_does_not_false_positive_on_real_project_name() -> None:
    """ชื่อโครงการจริงที่มีคำว่า "รวม" อยู่กลาง/ท้ายชื่อ (ไม่ใช่คำขึ้นต้น) ต้องไม่ถูกตัด"""
    header = ("ลำดับ", "ชื่อโครงการ", "หน่วยงาน", "งบประมาณ")
    mapping = map_columns(header)
    real_item_name = "โครงการระบบสูบน้ำ กลุ่มวิสาหกิจชุมชนรวมเกษตรยั่งยืน"
    real_row = (3, real_item_name, "เทศบาลตำบลบ้านแก้ง", 423000)
    assert _is_subtotal_row(real_item_name, real_row, mapping["agency"]) is False

    subtotal_row = (None, "รวม 2 แผน", None, 1_118_647_810)
    assert _is_subtotal_row("รวม 2 แผน", subtotal_row, mapping["agency"]) is True

    subtotal_row2 = (None, "ยอดรวมทั้งสิ้น", None, 1_000)
    assert _is_subtotal_row("ยอดรวมทั้งสิ้น", subtotal_row2, mapping["agency"]) is True


# ---------------------------------------------------------------------------
# gov_level / agency fallback
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("agency", "expected"),
    [
        ("เทศบาลตำบลบ้านโป่ง", "local"),
        ("องค์การบริหารส่วนตำบลป่าซาง", "local"),
        ("องค์การบริหารส่วนจังหวัดเชียงใหม่", "local"),
        ("กรุงเทพมหานคร", "local"),
        ("เมืองพัทยา", "local"),
        ("กรมชลประทาน", "central"),
        ("มหาวิทยาลัยเชียงใหม่", "central"),
        ("องค์การโคนม", "state_enterprise"),
        ("องค์การคลังสินค้า", "state_enterprise"),
    ],
)
def test_gov_level_for_agency(agency: str, expected: str) -> None:
    assert _gov_level_for_agency(agency, "") == expected


def test_agency_from_rel_path_skips_numbered_data_subfolder() -> None:
    """`1. ข้อมูลโครงการ` เป็น subfolder จัดหมวดข้อมูล ไม่ใช่ชื่อหน่วยงาน (พบจริงใน
    กองทุนอนุรักษ์พลังงาน) — ต้องข้ามไปใช้ folder ก่อนหน้าแทน
    """
    rel_path = (
        "กมธ.ติดตามงบ/ครั้งที่ 12 (13 ส.ค. 2569)/กองทุนอนุรักษ์พลังงาน/1. ข้อมูลโครงการ/รวมข้อมูลโครงการ.xlsx"
    )
    assert _agency_from_rel_path(rel_path) == "กองทุนอนุรักษ์พลังงาน"


def test_agency_from_rel_path_normal_two_level() -> None:
    rel_path = "กมธ.ติดตามงบ/ครั้งที่ 3 (28 พ.ค. 2569)/องค์การโคนม/BIS65R~1.XLS"
    assert _agency_from_rel_path(rel_path) == "องค์การโคนม"


# ---------------------------------------------------------------------------
# fiscal year extraction — ลำดับความสำคัญ sheet name ก่อนเสมอ + ตัด "ณ วันที่" ทิ้ง
# ---------------------------------------------------------------------------


def test_fiscal_year_from_sheet_name_with_prefix() -> None:
    assert _extract_fiscal_year_be("ปี 63", []) == 2563
    assert _extract_fiscal_year_be("ปี 64 กลุ่ม 7", []) == 2564


def test_fiscal_year_from_sheet_name_bare_leading_digits() -> None:
    """`61เพิ่มเติม` ไม่มีคำ "ปี" นำ แต่ขึ้นต้นด้วยเลข 2 หลักตรง ๆ (พบจริงใน กองทุนอนุรักษ์พลังงาน)"""
    assert _extract_fiscal_year_be("61เพิ่มเติม", []) == 2561
    assert _extract_fiscal_year_be("61เพิ่มเติม-รอบ2", []) == 2561


def test_fiscal_year_sheet_name_wins_over_conflicting_as_of_date_in_title() -> None:
    """บั๊กจริงที่แก้ 19 ก.ย. 2569: ไตเติลมีทั้งปีงบจริงกับวันที่ปรับปรุงข้อมูลปนกัน — ต้องไม่กำกวม
    เพราะชื่อ sheet ให้คำตอบตรง ๆ อยู่แล้ว
    """
    texts = [
        "ปี 63",
        "รวมข้อมูลโครงการ",
        "รายละเอียดโครงการที่อนุมัติในปีงบประมาณ พ.ศ. 2563",
        "ข้อมูล ณ วันที่ 30 มิถุนายน 2567",
    ]
    assert _extract_fiscal_year_be("ปี 63", texts) == 2563


def test_fiscal_year_fallback_to_title_excludes_as_of_date_text() -> None:
    """sheet name ไม่มีปี → fallback ไป title/header แต่ต้องตัดบรรทัด "ณ วันที่ ..." ทิ้งก่อน"""
    texts = [
        "Sheet1",
        "stem",
        "รายละเอียดโครงการที่อนุมัติในปีงบประมาณ พ.ศ. 2565",
        "ข้อมูล ณ วันที่ 1 มกราคม 2569",
    ]
    assert _extract_fiscal_year_be("Sheet1", texts) == 2565


def test_fiscal_year_ambiguous_returns_none() -> None:
    texts = ["Sheet2", "stem", "เทียบปีงบประมาณ พ.ศ. 2563 กับ พ.ศ. 2564"]
    assert _extract_fiscal_year_be("Sheet2", texts) is None


# ---------------------------------------------------------------------------
# `.XLS` ที่เนื้อไฟล์เป็น xlsx จริง (magic bytes = zip) ต้องเปิดได้ด้วย openpyxl
# ---------------------------------------------------------------------------


def test_xlsx_content_with_xls_extension_opens_via_openpyxl(tmp_path: Path) -> None:
    """regression: `องค์การโคนม/BI*.XLS` ทั้ง 13 ไฟล์ — เนื้อไฟล์เป็น xlsx (zip) แต่ตั้งนามสกุล
    `.XLS` เดิม `openpyxl.load_workbook(str(path), ...)` raise "does not support the old .xls
    file format" เพราะเช็คนามสกุลจาก path string เอง (ไม่ใช่แค่ magic bytes) แม้ routing จะ
    เลือก reader ถูกแล้วก็ตาม — ต้องเปิดด้วย file object (`open(path, "rb")`) เสมอ
    """
    path = tmp_path / "BI1C2F~1.XLS"
    _write_workbook(path, {"sheet1": [["a", "b"], [1, 2]]})

    assert _is_zip_magic(path) is True  # เนื้อไฟล์เป็น zip/xlsx จริง แม้นามสกุลเป็น .XLS

    sheets = dict(_iter_workbook_sheets(path))
    assert sheets["sheet1"] == [("a", "b"), (1, 2)]


def test_genuine_broken_xls_ole_header_raises_unsupported(tmp_path: Path) -> None:
    """ไฟล์ `.xls` แท้ (OLE magic `D0CF11E0`) ที่เปิดไม่ได้จริง (เสีย/ไม่ใช่ workbook) → skip
    ด้วยเหตุผลชัดเจน ไม่ raise exception ดิบออกไป
    """
    path = tmp_path / "broken.xls"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"not a real xls workbook body")

    assert _is_zip_magic(path) is False
    with pytest.raises(UnsupportedWorkbookError):
        _iter_workbook_sheets(path)


# ---------------------------------------------------------------------------
# discover_committee_files
# ---------------------------------------------------------------------------


def test_discover_committee_files_scoped_to_committee_root_and_sorted(tmp_path: Path) -> None:
    cfg = PipelineConfig.load(_write_config(tmp_path))
    committee_dir = cfg.raw_data_dir / COMMITTEE_ROOT / "ครั้งที่ 1" / "หน่วยงาน B"
    committee_dir.mkdir(parents=True)
    other_dir = cfg.raw_data_dir / "PBO"
    other_dir.mkdir(parents=True)

    _write_workbook(committee_dir / "z_file.xlsx", {"s": [["a"]]})
    _write_workbook(committee_dir / "a_file.xlsx", {"s": [["a"]]})
    _write_workbook(other_dir / "2566.xlsx", {"s": [["a"]]})  # นอก กมธ.ติดตามงบ — ต้องไม่ถูกนับ
    (committee_dir / "not_a_workbook.txt").write_text("x", encoding="utf-8")

    found = discover_committee_files(cfg)
    names = [p.name for p in found]
    assert names == ["a_file.xlsx", "z_file.xlsx"]


# ---------------------------------------------------------------------------
# extract_committee_file — end-to-end (parquet + doc_chunk)
# ---------------------------------------------------------------------------


def test_extract_committee_file_end_to_end(tmp_path: Path) -> None:
    cfg = PipelineConfig.load(_write_config(tmp_path))
    committee_dir = cfg.raw_data_dir / COMMITTEE_ROOT / "ครั้งที่ 1 (1 ม.ค. 2569)" / "กองทุนทดสอบ"
    committee_dir.mkdir(parents=True)
    path = committee_dir / "รวมข้อมูลโครงการ.xlsx"

    unmapped_rows = [
        ["ข้อมูลระบบ BIS", None, None, None],
        ["ลำดับ", "รายการ", "หน่วยงาน", "จำนวนเงิน"],
        [1, "รายการทดสอบ", "กรมทดสอบ", 5000],
    ]
    _write_workbook(path, {"ปี 63": _mapped_sheet_rows(), "mask1": unmapped_rows})

    rel_path = path.relative_to(cfg.raw_data_dir).as_posix()
    doc_id = doc_id_for_path(rel_path)

    result = extract_committee_file(cfg, path)

    assert result.doc_id == doc_id
    assert result.skipped is False
    assert result.mapped_rows == 2
    assert result.subtotal_rows_skipped == 1
    assert result.sheets_mapped == ["ปี 63"]
    assert result.sheets_unmapped == [("mask1", "ไม่รู้หน่วยเงิน (ไม่พบ บาท/พันบาท/ล้านบาท)")]

    assert result.parquet_path is not None
    table = pq.read_table(str(result.parquet_path))
    records = table.to_pylist()
    assert len(records) == 2
    assert {r["dataset"] for r in records} == {DATASET}
    assert {r["source_doc_id"] for r in records} == {doc_id}
    row_fay = next(r for r in records if r["agency"] == "เทศบาลตำบลบ้านโป่ง")
    expected_id = compute_source_id(DATASET, rel_path, "ปี 63", row_fay["source_row"])
    assert row_fay["source_id"] == expected_id
    assert row_fay["gov_level"] == "local"
    assert row_fay["fiscal_year_be"] == 2563

    assert result.doc_chunk_path is not None
    chunks = read_doc_chunks_gz(result.doc_chunk_path)
    assert len(chunks) >= 1
    assert all(c["doc_id"] == doc_id for c in chunks)
    joined_text = "\n".join(c["text"] for c in chunks)
    assert "รายการทดสอบ" in joined_text
    tables_with_sheet = [t for c in chunks for t in c["tables"] if t.get("sheet") == "mask1"]
    assert tables_with_sheet  # ตารางของ sheet ที่ map ไม่ได้ต้องแนบชื่อ sheet ไว้


def test_extract_committee_file_source_id_stable_across_runs(tmp_path: Path) -> None:
    cfg = PipelineConfig.load(_write_config(tmp_path))
    committee_dir = cfg.raw_data_dir / COMMITTEE_ROOT / "ครั้งที่ 1" / "กองทุนทดสอบ"
    committee_dir.mkdir(parents=True)
    path = committee_dir / "ไฟล์.xlsx"
    _write_workbook(path, {"ปี 63": _mapped_sheet_rows()})

    result1 = extract_committee_file(cfg, path)
    ids1 = sorted(r["source_id"] for r in pq.read_table(str(result1.parquet_path)).to_pylist())

    result2 = extract_committee_file(cfg, path)
    ids2 = sorted(r["source_id"] for r in pq.read_table(str(result2.parquet_path)).to_pylist())

    assert ids1 == ids2
    assert len(ids1) == len(set(ids1))  # unique ต่อไฟล์ (V4)
